"""Книга Excel с графиками (ТЗ, раздел 8; образец - листы PDF).

Листы книги
-----------
- лист данных - значения исходного листа без изменений (копируются потоком из файла, а
  не из ``Survey``: ничего не теряется, включая нераспознанные колонки и GPS). Называется
  так же, как исходный лист: ``Template`` или ``Result`` (ТЗ, раздел 8);
- «Карта» - на каждый участок две страницы печати (как листы 1 и 2 PDF): графики Excel
  (точечные диаграммы с прямыми линиями) и картинки лент «Тип изоляции», «Тип грунта»,
  «Объекты на трассе»;
- «Повреждения» - таблица участков повреждений ИП (ТЗ, 5.3);
- «Нормативы» - служебный лист: начальный километраж, две точки для линий нормативов и
  зон, ломаная «Площадь дефекта» по каждому участку;
- «Километраж» - служебный лист: километраж (``start_km + дистанция / 1000``) для осей X,
  строка в строку с листом данных (лист данных не изменяется).

Серии графиков ссылаются на диапазоны этих листов в пределах строк участка, данные не
дублируются: график правится и пересчитывается вместе с данными. Диапазон строк участка
для каждого ряда находится по колонке дистанции его блока (у грунта и глубины она своя).

Книга пишется потоково (``Workbook(write_only=True)``): этот режим поддерживает и
диаграммы, и картинки. Картинки лент пишутся во временную папку и читаются openpyxl
только при сохранении, поэтому память не растёт с числом участков.
"""

from __future__ import annotations

import math
import tempfile
import time
from array import array
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
from matplotlib.colors import to_hex
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.image import Image as XlImage
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.text import (
    CharacterProperties,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
    RichTextProperties,
)
from openpyxl.drawing.text import Font as DrawingFont
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from cormap.defects import DefectSegment, strip_steps
from cormap.layout import Section, format_km, place_objects
from cormap.model import Survey
from cormap.notes import TrackObject
from cormap.params import BuildParams
from cormap.reader import _EMPTY, _OK, _header_map, _parse_number
from cormap.render_pdf import (
    DEFECTS_NOTE,
    IMAGE_PAD_MM,
    STRIP_INSULATION,
    STRIP_OBJECTS,
    STRIP_SOIL,
    TITLE_DEFECTS,
    TITLE_DEPTH,
    TITLE_GRADIENT,
    TITLE_POTENTIAL,
    TITLE_RESISTIVITY,
    TITLE_ZONES,
    ZONES_NOTE,
    SheetContext,
    depth_axis,
    gradient_axis,
    potential_axis,
    render_strip_png,
    resistivity_axis,
    sheet1_heights,
    sheet2_heights,
    strip_image_size_mm,
    window,
    wrap_text,
)

__all__ = [
    "SHEET_DEFECTS",
    "SHEET_KM",
    "SHEET_MAP",
    "SHEET_NORMS",
    "write_xlsx",
]

SHEET_MAP = "Карта"
SHEET_DEFECTS = "Повреждения"
SHEET_NORMS = "Нормативы"
SHEET_KM = "Километраж"

NO_DATA = "нет данных"
NO_DEFECTS = "Повреждений нет"
IMAGE_DPI = 150

MM_EMU = 36000  # EMU в миллиметре
ROW_PT = 15.0  # высота строки листа «Карта» (фиксирована)
ROW_MM = ROW_PT * 25.4 / 72.0
ROW_EMU = round(ROW_PT * 12700)
MAP_COLUMNS = 12  # число колонок листа «Карта» (вместе они дают ширину рамки полосы)
HEADER_ROWS = 2  # заголовок участка на странице: объект, затем «км от - до, Участок N из M»
MIN_PAGE_MARGIN_MM = 8.0  # наименьшие поля сверху и снизу
PAPER_CODES = {"A4": 9, "A3": 8}  # коды paperSize в Excel

# Ряды графиков: ключ -> (ключ колонки дистанции, ключ колонки значений)
_SERIES: dict[str, tuple[str, str]] = {
    "u_on": ("dist", "u_on"),
    "u_off": ("dist", "u_off"),
    "gl_on": ("dist", "gl_on"),
    "gl_off": ("dist", "gl_off"),
    "gr_on": ("dist", "gr_on"),
    "gr_off": ("dist", "gr_off"),
    "rho": ("soil_dist", "rho"),
    "depth": ("depth_dist", "depth"),
}
# Колонки листа «Километраж» (номера с 1): по одной на колонку дистанции
_KM_COLUMNS = {"dist": 1, "soil_dist": 2, "depth_dist": 3}

StageTimer = Callable[[str, float], float]


# --- чтение исходного листа ---------------------------------------------------


@dataclass(slots=True)
class _Rows:
    """Строки листа данных, где у ряда есть дистанция и значение (по возрастанию строки)."""

    rows: array = field(default_factory=lambda: array("q"))
    dist: array = field(default_factory=lambda: array("d"))


@dataclass(slots=True)
class _Scan:
    """Итог прохода по исходному листу."""

    n_rows: int = 0
    n_cols: int = 0
    cols: dict[str, int] = field(default_factory=dict)  # ключ колонки -> индекс с 0
    headers: tuple[Any, ...] = ()
    series: dict[str, _Rows] = field(default_factory=dict)
    defect_rows: dict[float, list[int]] = field(default_factory=dict)  # дистанция -> строки


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool) and math.isfinite(value)


def _copy_source(
    src: str | Path,
    sheet_name: str,
    ws_data: Any,
    ws_km: Any,
    start_km: float,
) -> _Scan:
    """Копирует исходный лист потоком в лист данных, строит «Километраж», ищет строки рядов."""
    scan = _Scan()
    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        ws.reset_dimensions()
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, ())
        cols, _ = _header_map(header)
        scan.cols = cols
        scan.headers = header
        scan.n_cols = len(header)
        km_idx = {dk: cols[dk] for dk in _KM_COLUMNS if dk in cols}
        series = {k: v for k, v in _SERIES.items() if v[0] in cols and v[1] in cols}
        scan.series = {k: _Rows() for k in series}
        dmg_ok = "dmg_dist" in cols and "dmg" in cols
        letters = {dk: _col_letter(i + 1) for dk, i in km_idx.items()}
        ws_data.append(_safe_row(ws_data, header))
        ws_km.append([f"км: {header[km_idx[dk]]}" if dk in km_idx else None for dk in _KM_COLUMNS])
        scan.n_rows = 1
        norms = f"'{SHEET_NORMS}'!$B$1"
        for rownum, row in enumerate(rows_iter, 2):
            ws_data.append(_safe_row(ws_data, row))
            scan.n_rows = rownum
            n = len(row)
            dists: dict[str, float] = {}
            km_cells: list[Any] = [None, None, None]
            for dk, idx in km_idx.items():
                value = row[idx] if idx < n else None
                if _is_number(value):
                    dists[dk] = float(value)
                    km_cells[_KM_COLUMNS[dk] - 1] = (
                        f"='{ws_data.title}'!{letters[dk]}{rownum}/1000+{norms}"
                    )
                else:
                    status, number = _parse_number(value)
                    if status == _OK:
                        dists[dk] = number
                        km_cells[_KM_COLUMNS[dk] - 1] = start_km + number / 1000.0
            ws_km.append(km_cells)
            if dmg_ok:
                di = cols["dmg_dist"]
                status, number = _parse_number(row[di] if di < n else None)
                if status == _OK:
                    dists["dmg_dist"] = number
            if not dists:
                continue
            for key, (dk, vk) in series.items():
                d = dists.get(dk)
                if d is None:
                    continue
                vi = cols[vk]
                status, _ = _parse_number(row[vi] if vi < n else None)
                if status == _OK:
                    s = scan.series[key]
                    s.rows.append(rownum)
                    s.dist.append(d)
            if dmg_ok and "dmg_dist" in dists:
                vi = cols["dmg"]
                raw = row[vi] if vi < n else None
                if _parse_number(raw)[0] != _EMPTY:
                    scan.defect_rows.setdefault(dists["dmg_dist"], []).append(rownum)
    finally:
        wb.close()
    return scan


def _col_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _safe_row(ws: Any, row: Sequence[Any]) -> list[Any]:
    """Строка для записи: текст, начинающийся с «=», остаётся текстом, а не формулой."""
    if not any(isinstance(v, str) and v[:1] == "=" for v in row):
        return list(row)
    out: list[Any] = []
    for v in row:
        if isinstance(v, str) and v[:1] == "=":
            cell = WriteOnlyCell(ws, value=v)
            cell.data_type = "s"
            out.append(cell)
        else:
            out.append(v)
    return out


# --- диапазоны строк участка ---------------------------------------------------


@dataclass(frozen=True, slots=True)
class _Range:
    """Строки листа «Данные» для ряда на участке (включая по одной соседней точке)."""

    first: int
    last: int
    has_data: bool  # есть что рисовать: значения внутри участка или линия сквозь участок


class _RowIndex:
    """Поиск строк участка по дистанции для каждого ряда."""

    def __init__(self, scan: _Scan) -> None:
        self._rows = {k: np.frombuffer(v.rows, dtype=np.int64) for k, v in scan.series.items()}
        self._dist = {k: np.frombuffer(v.dist, dtype=np.float64) for k, v in scan.series.items()}

    def find(self, key: str, a: float, b: float) -> _Range | None:
        """Диапазон строк ряда для участка [a, b]; None - ряда нет или нет строк."""
        dist = self._dist.get(key)
        if dist is None or dist.size == 0:
            return None
        rows = self._rows[key]
        i0 = int(np.searchsorted(dist, a, side="left"))
        i1 = int(np.searchsorted(dist, b, side="right"))
        inside = i1 - i0
        has = inside > 0 or (i0 > 0 and i1 < dist.size)
        if not has:
            return None
        lo = max(i0 - 1, 0)
        hi = min(i1, dist.size - 1)
        return _Range(int(rows[lo]), int(rows[hi]), True)


# --- оформление диаграмм -------------------------------------------------------


def _hex(color: str) -> str:
    return to_hex(color, keep_alpha=False).lstrip("#").upper()


def _char_props(size_pt: float, family: str, *, bold: bool = False) -> CharacterProperties:
    return CharacterProperties(
        sz=round(size_pt * 100), b=bold, latin=DrawingFont(typeface=family), solidFill="000000"
    )


def _txpr(size_pt: float, family: str, *, bold: bool = False) -> RichText:
    cp = _char_props(size_pt, family, bold=bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _rich_title(
    lines: Sequence[str],
    size_pt: float,
    family: str,
    *,
    bold: bool,
    rotate: bool = False,
    layout: Layout | None = None,
    overlay: bool = False,
) -> Title:
    cp = _char_props(size_pt, family, bold=bold)
    paragraphs = [
        Paragraph(
            pPr=ParagraphProperties(algn="l" if layout is not None else "ctr", defRPr=cp),
            r=[RegularTextRun(rPr=cp, t=line)],
        )
        for line in lines
    ]
    body = RichTextProperties(rot=-5400000, vert="horz") if rotate else RichTextProperties()
    return Title(tx=Text(rich=RichText(bodyPr=body, p=paragraphs)), layout=layout, overlay=overlay)


@dataclass(slots=True)
class _Line:
    """Серия графика: подпись, ссылки на ячейки, вид линии."""

    title: str
    x: Reference
    y: Reference
    color: str
    width_pt: float
    dash: bool = False


@dataclass(frozen=True, slots=True)
class _Axis:
    lo: float
    hi: float
    step: float
    reverse: bool = False


class _Maker:
    """Собирает диаграммы одного вида и размера: общая геометрия, шрифты, цвета."""

    def __init__(self, params: BuildParams, ctx: SheetContext) -> None:
        self.params = params
        self.ctx = ctx
        self.ls, self.rs = ctx.layout, ctx.render
        self.family = self.ls.font_families[0]
        paper = self.ls.paper(params.paper)
        self.width_mm = paper.width_mm - 2 * self.ls.margin_mm + 2 * IMAGE_PAD_MM
        self.plot_left_mm = IMAGE_PAD_MM + self.ls.label_column_mm + self.ls.plot_pad_left_mm
        self.plot_w_mm = self.ls.plot_width_mm(params.paper)

    def chart(
        self,
        title: str,
        height_mm: float,
        lines: Sequence[_Line],
        x_axis: _Axis,
        y_axis: _Axis,
        y_title: str,
        *,
        grid_y: bool = True,
    ) -> ScatterChart:
        rs, g = self.rs, self.rs.geometry
        w, h = self.width_mm, height_mm
        ch = ScatterChart()
        ch.scatterStyle = "lineMarker"
        ch.style = 2
        ch.roundedCorners = False
        ch.width, ch.height = w / 10.0, h / 10.0
        ch.display_blanks = "span"
        ch.varyColors = False
        # название полосы - в левой колонке, как на листе PDF (с переносом по словам)
        pt = rs.sizes.strip_label_pt
        wrapped = wrap_text(title, pt, self.ls.label_column_mm - 4.0, tuple(self.ls.font_families))
        ch.title = _rich_title(
            wrapped,
            pt,
            self.family,
            bold=True,
            overlay=True,
            layout=Layout(
                manualLayout=ManualLayout(
                    xMode="edge", yMode="edge", x=(IMAGE_PAD_MM + 1.5) / w, y=0.28
                )
            ),
        )
        # область графика: тот же отступ слева и ширина, что на листе PDF
        top_pad = g.axes_top_pad_mm
        bottom = g.tick_row_mm + 4.0 + g.legend_box_mm + 2 * g.legend_pad_mm
        ch.layout = Layout(
            manualLayout=ManualLayout(
                layoutTarget="inner",
                xMode="edge",
                yMode="edge",
                x=self.plot_left_mm / w,
                y=top_pad / h,
                w=self.plot_w_mm / w,
                h=max(h - top_pad - bottom, 4.0) / h,
            )
        )
        ch.legend.position = "b"
        ch.legend.overlay = False
        ch.legend.txPr = _txpr(rs.sizes.legend_pt, self.family)
        tick_pt = rs.sizes.tick_pt
        axis_line = LineProperties(solidFill=_hex(rs.colors.axis), w=round(rs.lines.axis * 12700))
        grid_line = LineProperties(solidFill="BFBFBF", w=round(rs.lines.grid * 12700))
        for axis, spec, fmt in ((ch.x_axis, x_axis, "0.00"), (ch.y_axis, y_axis, "0.00")):
            axis.delete = False
            axis.scaling.min, axis.scaling.max = spec.lo, spec.hi
            axis.majorUnit = spec.step
            axis.number_format = fmt
            axis.txPr = _txpr(tick_pt, self.family)
            axis.spPr = GraphicalProperties(ln=axis_line)
            axis.majorTickMark = "out"
            axis.minorTickMark = "none"
        ch.y_axis.number_format = "0.00" if (y_axis.step % 1) else "0"
        ch.x_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=grid_line))
        ch.y_axis.majorGridlines = (
            ChartLines(spPr=GraphicalProperties(ln=grid_line)) if grid_y else None
        )
        ch.x_axis.title = _rich_title(
            ["Протяженность, км"], rs.sizes.caption_pt, self.family, bold=False
        )
        ch.y_axis.title = _rich_title(
            [y_title], rs.sizes.caption_pt, self.family, bold=False, rotate=True
        )
        ch.y_axis.crosses = "min"
        if y_axis.reverse:
            ch.y_axis.scaling.orientation = "maxMin"
            ch.x_axis.crosses = "max"  # ось X остаётся внизу при перевёрнутой оси Y
        for line in lines:
            s = Series(line.y, line.x, title=line.title)
            s.marker = Marker(symbol="none")
            s.smooth = False
            gp = GraphicalProperties()
            gp.line = LineProperties(
                solidFill=_hex(line.color),
                w=round(max(line.width_pt, 0.5) * 12700),
                prstDash="dash" if line.dash else "solid",
            )
            s.graphicalProperties = gp
            ch.series.append(s)
        return ch


# --- служебные листы -------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _NormRows:
    """Строки листа «Нормативы» для участка."""

    pair: int  # первая из двух строк (начало и конец оси)
    steps_first: int  # ломаная «Площадь дефекта»: первая и последняя строки
    steps_last: int


def _defect_polyline(
    defects: Sequence[DefectSegment], start: float, end: float
) -> list[tuple[float, int]]:
    """Ломаная «Площадь дефекта» участка (дистанция, класс): ступеньки и точечные штрихи."""
    items: list[tuple[float, float, int]] = list(strip_steps(defects, start, end))
    for seg in defects:  # точечные повреждения с классом - вертикальный штрих (как в PDF)
        if seg.class_final and seg.start == seg.end and start <= seg.start <= end:
            items.append((seg.start, seg.start, seg.class_final))
    items.sort()
    points: list[tuple[float, int]] = [(start, 0)]
    for lo, hi, cls in items:
        points += [(lo, 0), (lo, cls), (hi, cls), (hi, 0)]
    points.append((end, 0))
    out: list[tuple[float, int]] = []
    for pt in points:  # соседние одинаковые точки не нужны
        if not out or out[-1] != pt:
            out.append(pt)
    return out


def _write_norms(
    ws: Any,
    sections: Sequence[Section],
    params: BuildParams,
    ctx: SheetContext,
) -> dict[int, _NormRows]:
    """Лист «Нормативы». Возвращает строки по номеру участка."""
    bold = Font(bold=True)
    head_rows: list[list[Any]] = [
        [
            WriteOnlyCell(ws, "Начальный километраж, км"),
            params.start_km,
            None,
            WriteOnlyCell(
                ws, "Служебный лист: значения линий нормативов и зон на графиках листа «Карта»."
            ),
        ],
        [],
    ]
    head = [
        "Участок",
        "Дистанция, м",
        "Километраж, км",
        "U т-з макс. норматив, В",
        "U т-з мин. норматив, В",
        "Мин. глубина, м",
        "Порог УЭС нижний, Ом·м",
        "Порог УЭС верхний, Ом·м",
        "Зоны (0 - УКО, 1 - ВКО): линия на нуле",
        None,
        "Участок",
        "Дистанция, м",
        "Километраж, км",
        "Класс (площадь дефекта)",
    ]
    cells = []
    for text in head:
        if text is None:
            cells.append(None)
        else:
            c = WriteOnlyCell(ws, text)
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical="top")
            cells.append(c)
    # левая таблица: две строки на участок; правая - ломаная дефектов (строк по-разному)
    left: list[list[Any]] = []
    right: list[list[Any]] = []
    refs: dict[int, _NormRows] = {}
    first_data = len(head_rows) + 2  # строка заголовка таблицы - 3, данные с 4
    for sec in sections:
        pair = first_data + len(left)
        for dist in (sec.start_m, sec.start_m + sec.nominal_length_m):
            r = first_data + len(left)
            left.append(
                [
                    sec.index,
                    dist,
                    f"=$B$1+B{r}/1000",
                    params.u_max,
                    params.u_min,
                    params.depth_min,
                    params.res_low,
                    params.res_high,
                    0,
                ]
            )
        steps_first = first_data + len(right)
        for dist, cls in _defect_polyline(ctx.defects, sec.start_m, sec.end_m):
            r = first_data + len(right)
            right.append([sec.index, dist, f"=$B$1+L{r}/1000", cls])
        refs[sec.index] = _NormRows(pair, steps_first, first_data + len(right) - 1)
    for row in head_rows:
        ws.append(row)
    ws.append(cells)
    for i in range(max(len(left), len(right))):
        a = left[i] if i < len(left) else [None] * 9
        b = right[i] if i < len(right) else [None] * 4
        ws.append([*a, None, *b])
    return refs


def _write_defects(ws: Any, ctx: SheetContext, survey: Survey, scan: _Scan, start_km: float) -> int:
    """Лист «Повреждения» (ТЗ, 5.3). Возвращает число строк таблицы."""
    bold = Font(bold=True)
    head = [
        "Начало, м",
        "Конец, м",
        "Начало, км",
        "Конец, км",
        "Протяжённость, м",
        "Наибольшее отклонение",
        "Класс по силе",
        "Класс по протяжённости",
        "Итоговый класс",
        "Источник",
        "Строка книги",
    ]
    cells = []
    for text in head:
        c = WriteOnlyCell(ws, text)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
        cells.append(c)
    ws.append(cells)
    if not ctx.defects:
        ws.append([NO_DEFECTS])
        return 0
    block_rows = {d: list(r) for d, r in scan.defect_rows.items()}
    for seg in ctx.defects:
        row = seg.row
        if row is None and seg.source == "manual" and seg.start == seg.end:
            found = block_rows.get(seg.start)
            if found:
                row = found.pop(0)
        ws.append(
            [
                seg.start,
                seg.end,
                round(start_km + seg.start / 1000.0, 6),
                round(start_km + seg.end / 1000.0, 6),
                seg.length,
                seg.max_deviation,
                seg.class_strength,
                seg.class_length,
                seg.class_final,
                "расчёт" if seg.source == "calc" else "ручной",
                row,
            ]
        )
    del survey
    return len(ctx.defects)


# --- лист «Карта» ----------------------------------------------------------------


@dataclass(slots=True)
class _Page:
    """Страница печати участка: строки с текстом, диаграммы и картинки."""

    texts: dict[int, list[Any]] = field(default_factory=dict)  # строка страницы (с 0) -> ячейки


class _MapWriter:
    """Раскладка листа «Карта»: страницы одинаковой высоты, элементы сверху вниз."""

    def __init__(
        self,
        wb: Workbook,
        ws: Any,
        params: BuildParams,
        ctx: SheetContext,
        tmp: Path,
    ) -> None:
        self.ws, self.params, self.ctx, self.tmp = ws, params, ctx, tmp
        self.maker = _Maker(params, ctx)
        paper = ctx.layout.paper(params.paper)
        # строк на страницу: целое число, поля сверху и снизу не меньше минимума
        self.page_rows = int((paper.height_mm - 2 * MIN_PAGE_MARGIN_MM) // ROW_MM)
        self.margin_v_mm = (paper.height_mm - self.page_rows * ROW_MM) / 2
        self.margin_h_mm = ctx.layout.margin_mm - IMAGE_PAD_MM
        self.width_mm = self.maker.width_mm
        self.n_pages = 0
        self.image_no = 0
        wb.calculation.fullCalcOnLoad = True

    # --- настройка листа до записи строк ---

    def setup(self, n_sections: int) -> None:
        ws, p = self.ws, self.params
        paper = self.ctx.layout.paper(p.paper)
        px = self.width_mm / 25.4 * 96.0 / MAP_COLUMNS
        for i in range(MAP_COLUMNS):
            ws.column_dimensions[_col_letter(i + 1)].width = round((px - 5.0) / 7.0, 3)
        ws.sheet_format.defaultRowHeight = ROW_PT
        ws.sheet_format.customHeight = True
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = PAPER_CODES[paper.name.upper()]
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        inch = 25.4
        ws.page_margins.left = ws.page_margins.right = self.margin_h_mm / inch
        ws.page_margins.top = ws.page_margins.bottom = self.margin_v_mm / inch
        ws.page_margins.header = ws.page_margins.footer = 3.0 / inch
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "Лист &P из &N"
        ws.oddFooter.center.size = 8
        self.n_pages = 2 * n_sections
        ws.print_area = f"A1:{_col_letter(MAP_COLUMNS)}{self.n_pages * self.page_rows}"
        for k in range(1, self.n_pages):
            ws.row_breaks.append(Break(id=k * self.page_rows))

    # --- элементы ---

    def _anchor(self, page: int, y_mm: float, h_mm: float) -> OneCellAnchor:
        row, off = divmod(round(y_mm * MM_EMU), ROW_EMU)
        return OneCellAnchor(
            _from=AnchorMarker(col=0, colOff=0, row=page * self.page_rows + row, rowOff=off),
            ext=XDRPositiveSize2D(cx=round(self.width_mm * MM_EMU), cy=round(h_mm * MM_EMU)),
        )

    def _text(self, value: str, *, bold: bool, size: float, center: bool = False) -> list[Any]:
        cells: list[Any] = []
        for i in range(MAP_COLUMNS if center else 1):
            c = WriteOnlyCell(self.ws, value if i == 0 else None)
            c.font = Font(name=self.maker.family, bold=bold, size=size)
            if center:
                c.alignment = Alignment(horizontal="centerContinuous")
            cells.append(c)
        return cells

    def _header(self, page: _Page, sec: Section, number: int, of: int, sheet_no: int) -> None:
        p = self.params
        km0 = format_km(p.start_km + sec.start_m / 1000.0)
        km1 = format_km(p.start_km + sec.end_m / 1000.0)
        page.texts[0] = self._text(p.object_name, bold=True, size=12, center=True)
        page.texts[1] = self._text(
            f"км {km0} - {km1}.  Участок {number} из {of}.  Лист {sheet_no} из 2",
            bold=False,
            size=10,
            center=True,
        )

    def _no_data(self, page: _Page, y_mm: float, title: str) -> None:
        row = int(y_mm // ROW_MM) + 1
        page.texts[row] = self._text(f"{title}: {NO_DATA}", bold=True, size=10)

    def _flush(self, page: _Page) -> None:
        for r in range(self.page_rows):
            self.ws.append(page.texts.get(r, []))

    # --- участок ---

    def add_section(
        self,
        sec: Section,
        of: int,
        placement: Any,
        index: _RowIndex,
        scan: _Scan,
        norms: _NormRows,
        survey: Survey,
        km_sheet: Any,
        data_sheet: Any,
        norm_sheet: Any,
    ) -> None:
        ctx, p = self.ctx, self.params
        a, b = sec.start_m, sec.end_m
        pair = (norms.pair, norms.pair + 1)
        km0 = round(p.start_km + a / 1000.0, 6)
        km1 = round(p.start_km + (a + sec.nominal_length_m) / 1000.0, 6)
        x_axis = _Axis(km0, km1, sec.tick_step_m / 1000.0)
        d = ctx.data
        win = lambda s: window(s, a, b, 1 << 40)  # noqa: E731 - без прореживания: нужны границы
        mk = self.maker
        rs = ctx.render

        def ref(sheet: Any, col: int, r0: int, r1: int) -> Reference:
            return Reference(sheet, min_col=col, min_row=r0, max_row=r1)

        def data_line(title: str, key: str, color: str, width: float, rng: _Range) -> _Line:
            dk, vk = _SERIES[key]
            return _Line(
                title,
                ref(km_sheet, _KM_COLUMNS[dk], rng.first, rng.last),
                ref(data_sheet, scan.cols[vk] + 1, rng.first, rng.last),
                color,
                width,
            )

        def norm_line(title: str, col: int, color: str, width: float, dash: bool = True) -> _Line:
            return _Line(
                title,
                ref(norm_sheet, 3, *pair),
                ref(norm_sheet, col, *pair),
                color,
                width,
                dash,
            )

        # --- страница 1: потенциал, градиенты, площадь дефекта, зоны ---
        page1, page2 = _Page(), _Page()
        self._header(page1, sec, sec.index, of, 1)
        self._header(page2, sec, sec.index, of, 2)
        h1 = sheet1_heights(p, ctx)
        body = self.page_rows * ROW_MM - HEADER_ROWS * ROW_MM
        total1 = sum(h1.values()) + 5 * 2 * IMAGE_PAD_MM
        k1 = min(1.0, body / total1)
        y = HEADER_ROWS * ROW_MM
        page_no = 2 * (sec.index - 1)

        def place(page: int, y_mm: float, h_mm: float, obj: Any) -> None:
            obj.anchor = self._anchor(page, y_mm, h_mm)
            if isinstance(obj, ScatterChart):
                self.ws.add_chart(obj)
            else:
                self.ws.add_image(obj)

        # потенциал
        h = h1["potential"] * k1 + 2 * IMAGE_PAD_MM
        r_on, r_off = index.find("u_on", a, b), index.find("u_off", a, b)
        if r_on or r_off:
            c = rs.colors
            top, bottom, step = potential_axis(win(d.u_on), win(d.u_off), p, rs)
            lines = []
            if r_on:
                lines.append(
                    data_line("защитный потенциал", "u_on", c.u_on, rs.lines.potential, r_on)
                )
            if r_off:
                lines.append(
                    data_line("потенциал отключения", "u_off", c.u_off, rs.lines.potential, r_off)
                )
            lines.append(norm_line("максимальный нормативный потенциал", 4, c.u_max, rs.lines.norm))
            lines.append(norm_line("минимальный нормативный потенциал", 5, c.u_min, rs.lines.norm))
            place(
                page_no,
                y,
                h,
                mk.chart(TITLE_POTENTIAL, h, lines, x_axis, _Axis(top, bottom, step, True), "U, В"),
            )
        else:
            self._no_data(page1, y, TITLE_POTENTIAL)
        y += h
        # градиенты
        for side, word, name in (
            ("left", "левый", "gradient_left"),
            ("right", "правый", "gradient_right"),
        ):
            h = h1[name] * k1 + 2 * IMAGE_PAD_MM
            k_on, k_off = ("gl_on", "gl_off") if side == "left" else ("gr_on", "gr_off")
            r_on, r_off = index.find(k_on, a, b), index.find(k_off, a, b)
            if r_on or r_off:
                c_on = getattr(rs.colors, f"gradient_{side}_on")
                c_off = getattr(rs.colors, f"gradient_{side}_off")
                lo, hi, step = gradient_axis(
                    win(getattr(d, f"grad_{side}_on")), win(getattr(d, f"grad_{side}_off")), rs
                )
                lines = []
                if r_on:
                    lines.append(data_line(f"Градиент {word} вкл", k_on, c_on, 1.0, r_on))
                if r_off:
                    lines.append(data_line(f"Градиент {word} выкл", k_off, c_off, 1.0, r_off))
                place(
                    page_no,
                    y,
                    h,
                    mk.chart(
                        TITLE_GRADIENT, h, lines, x_axis, _Axis(lo, hi, step), "мВ", grid_y=False
                    ),
                )
            else:
                self._no_data(page1, y, f"{TITLE_GRADIENT} ({word})")
            y += h
        # площадь дефекта
        h = h1["defects"] * k1 + 2 * IMAGE_PAD_MM
        n = norms.steps_last - norms.steps_first + 1
        steps = _Line(
            DEFECTS_NOTE,
            ref(norm_sheet, 13, norms.steps_first, norms.steps_last),
            ref(norm_sheet, 14, norms.steps_first, norms.steps_last),
            rs.colors.defects,
            rs.lines.defects,
        )
        del n
        place(
            page_no,
            y,
            h,
            mk.chart(
                TITLE_DEFECTS,
                h,
                [steps],
                x_axis,
                _Axis(0.0, rs.axes.defects_top, 1.0),
                "Класс",
            ),
        )
        y += h
        # зоны коррозионной опасности: линия на нуле
        h = h1["zones"] * k1 + 2 * IMAGE_PAD_MM
        zero = norm_line(ZONES_NOTE, 9, rs.colors.frame, rs.lines.zones, dash=False)
        place(
            page_no,
            y,
            h,
            mk.chart(TITLE_ZONES, h, [zero], x_axis, _Axis(0.0, 1.0, 1.0), "Зона", grid_y=False),
        )
        self._flush(page1)

        # --- страница 2: изоляция, глубина, УЭС, грунт, объекты ---
        h2 = sheet2_heights(p, ctx, placement)
        heights = {
            "insulation": h2.insulation,
            "depth": h2.depth,
            "resistivity": h2.resistivity,
            "soil": h2.soil,
            "objects": h2.objects,
        }
        total2 = sum(heights.values()) + 5 * 2 * IMAGE_PAD_MM
        k2 = min(1.0, body / total2)
        y = HEADER_ROWS * ROW_MM
        page_no += 1

        def picture(kind: str, y_mm: float) -> float:
            png = render_strip_png(kind, sec, placement, p, ctx, dpi=IMAGE_DPI)
            w_mm, h_mm = strip_image_size_mm(kind, p, ctx, placement)
            h_mm = (h_mm - 2 * IMAGE_PAD_MM) * k2 + 2 * IMAGE_PAD_MM
            self.image_no += 1
            path = self.tmp / f"strip_{self.image_no}.png"
            path.write_bytes(png)
            img = XlImage(str(path))
            img.width, img.height = round(w_mm / 25.4 * 96), round(h_mm / 25.4 * 96)
            place(page_no, y_mm, h_mm, img)
            return h_mm

        y += picture(STRIP_INSULATION, y)
        # глубина
        h = heights["depth"] * k2 + 2 * IMAGE_PAD_MM
        r = index.find("depth", a, b)
        if r:
            lo, hi, step = depth_axis(win(d.depth), p, rs)
            lines = [
                data_line("глубина залегания", "depth", rs.colors.depth, rs.lines.depth, r),
                norm_line(
                    "минимальная нормативная глубина залегания",
                    6,
                    rs.colors.depth_norm,
                    rs.lines.depth_norm,
                ),
            ]
            place(
                page_no,
                y,
                h,
                mk.chart(TITLE_DEPTH, h, lines, x_axis, _Axis(lo, hi, step), "м"),
            )
        else:
            self._no_data(page2, y, TITLE_DEPTH)
        y += h
        # УЭС
        h = heights["resistivity"] * k2 + 2 * IMAGE_PAD_MM
        r = index.find("rho", a, b)
        if r:
            top, step = resistivity_axis(win(d.resistivity), p)
            c = rs.colors
            lines = [
                data_line("УЭС", "rho", c.resistivity, rs.lines.resistivity + 0.3, r),
                norm_line(f"порог УЭС {p.res_low:g} Ом·м", 7, c.res_high, rs.lines.norm),
                norm_line(f"порог УЭС {p.res_high:g} Ом·м", 8, c.res_low, rs.lines.norm),
            ]
            place(
                page_no,
                y,
                h,
                mk.chart(TITLE_RESISTIVITY, h, lines, x_axis, _Axis(0.0, top, step), "Ом·м"),
            )
        else:
            self._no_data(page2, y, TITLE_RESISTIVITY)
        y += h
        y += picture(STRIP_SOIL, y)
        picture(STRIP_OBJECTS, y)
        self._flush(page2)
        del survey


# --- сборка книги --------------------------------------------------------------


def write_xlsx(
    path: str | Path,
    source: str | Path,
    survey: Survey,
    sections: Sequence[Section],
    objects: Iterable[TrackObject],
    params: BuildParams,
    ctx: SheetContext,
    *,
    stage: StageTimer | None = None,
    progress: Callable[[int, int], None] | None = None,
) -> None:
    """Пишет книгу Excel ``path``: лист данных (копия листа ``source``), «Карта» и остальные.

    Лист данных называется как исходный лист книги (``survey.sheet_name``).

    ``objects`` - все объекты трассы (комментарии и ручные повреждения), как при разбиении.
    ``stage(имя, начало)`` - необязательный замер времени стадий; возвращает текущее время.
    """
    tick = stage or (lambda _name, _since: time.perf_counter())
    t = time.perf_counter()
    wb = Workbook(write_only=True)
    ws_data = wb.create_sheet(survey.sheet_name.strip())
    ws_map = wb.create_sheet(SHEET_MAP)
    ws_def = wb.create_sheet(SHEET_DEFECTS)
    ws_norm = wb.create_sheet(SHEET_NORMS)
    ws_km = wb.create_sheet(SHEET_KM)
    ws_data.freeze_panes = "A2"
    ws_km.freeze_panes = "A2"
    ws_norm.freeze_panes = "A4"
    ws_def.freeze_panes = "A2"
    def_widths = (12, 12, 12, 12, 16, 16, 14, 18, 14, 12, 12)
    for letter, width in zip("ABCDEFGHIJK", def_widths, strict=True):
        ws_def.column_dimensions[letter].width = width
    for letter in "ABCDEFGHI":
        ws_norm.column_dimensions[letter].width = 18
    for letter in "KLMN":
        ws_norm.column_dimensions[letter].width = 18
    for letter in "ABC":
        ws_km.column_dimensions[letter].width = 22
    ws_km.sheet_properties.tabColor = ws_norm.sheet_properties.tabColor = "A6A6A6"

    scan = _copy_source(source, survey.sheet_name, ws_data, ws_km, params.start_km)
    t = tick("Excel: лист данных", t)
    _write_defects(ws_def, ctx, survey, scan, params.start_km)
    norms = _write_norms(ws_norm, sections, params, ctx)
    index = _RowIndex(scan)
    t = tick("Excel: служебные листы", t)

    objs = list(objects)
    with tempfile.TemporaryDirectory(prefix="cormap_xlsx_") as tmp:
        writer = _MapWriter(wb, ws_map, params, ctx, Path(tmp))
        writer.setup(len(sections))
        for n, sec in enumerate(sections, 1):
            placement = place_objects(objs, sec, ctx.layout, params.paper)
            writer.add_section(
                sec,
                len(sections),
                placement,
                index,
                scan,
                norms[sec.index],
                survey,
                ws_km,
                ws_data,
                ws_norm,
            )
            if progress:
                progress(n, len(sections))
        t = tick("Excel: графики и картинки", t)
        wb.save(path)
        tick("Excel: запись файла", t)
