"""Сборка шаблона книги Excel (ТЗ, раздел 4.6).

Шаблон собирается из тех же словарей, по которым работает разбор: список объектов берётся
из `cormap.notes.OBJECT_TYPES`, допустимые типы грунта и изоляции - из `cormap.model`.
Модуль не зависит от Django.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from cormap import assets
from cormap.model import INSULATION_TYPES, SOIL_TYPES
from cormap.notes import (
    ABOVE,
    AREA_BELOW,
    BELOW,
    CROSS_BAND,
    CROSS_LINE,
    END,
    LEGACY,
    LEGEND,
    OBJECT_TYPES,
    ON,
    POINT,
    START,
    ObjectType,
)
from cormap.reader import COLUMNS

# Порядок колонок листа данных (ТЗ, 4.5): как в reference/template.xlsx, «Тип грунта» - сразу
# за УЭС, колонки изоляции - в конце. Модуль обработки находит колонки по номеру.
TEMPLATE_KEYS: tuple[str, ...] = (
    "dist",
    "u_on",
    "u_off",
    "gl_on",
    "gl_off",
    "gr_on",
    "gr_off",
    "comment",
    "gps",
    "soil_dist",
    "rho",
    "soil_type",
    "depth_dist",
    "depth",
    "dmg_dist",
    "dmg",
    "ins_dist",
    "ins_type",
)

SHEET_OBJECTS = "Объекты"
SHEET_HOWTO = "Как заполнять"
SHEET_LISTS = "Списки"

LAST_ROW = 10000  # проверка данных действует на строки 2...LAST_ROW

COMMENT_PROMPT_TITLE = "Объект на трассе"
COMMENT_PROMPT = (
    "Выберите объект из списка. После шифра можно дописать подпись через запятую: "
    "кип, КИП-14. Несколько объектов - через точку с запятой."
)

# Образцы словаря, записанные основой слова: для списка нужна полная запись.
FULL_FORMS: dict[str, str] = {"болот": "болото"}

PLACEMENT_TITLES: dict[str, str] = {
    ABOVE: "над трубой",
    ON: "на трубе",
    BELOW: "под трубой",
    CROSS_LINE: "линия поперёк трубы",
    CROSS_BAND: "полоса поперёк трубы",
    AREA_BELOW: "участок под трубой",
}

ROLE_SUFFIX = {START: " начало", END: " конец"}

HOWTO_LINES: tuple[str, ...] = (
    "Как заполнять книгу",
    "",
    "1. Данные вносятся на первом листе, заголовки в первой строке, данные со второй.",
    "2. Дистанции указываются в метрах от начала обследованного участка.",
    "3. У каждого блока своя колонка дистанции (интенсивные измерения, грунт, глубина, "
    "повреждения, изоляция): шаги блоков друг от друга не зависят.",
    "4. В колонке «Комментарий» одна пометка записывается как «тип, подпись». Подпись "
    "необязательна. Пример: кип, КИП-14.",
    "5. Тип выбирается из выпадающего списка (см. лист «Объекты»); подпись дописывается "
    "после запятой.",
    "6. Несколько пометок в одной ячейке разделяются точкой с запятой: укз, УКЗ-3; кип.",
    "7. Протяжённый объект задаётся парой строк: «река начало, р. Упа» и ниже «река конец». "
    "Пометка без слов «начало» и «конец» рисуется как точечный объект.",
    "8. Начало и конец должны быть у одного и того же типа объекта; одиночное «начало» "
    "или «конец» попадёт в отчёт проверки как ошибка.",
    "9. Для высоковольтной линии напряжение пишется через дефис: ВЛ-110 кВ.",
    "10. Регистр букв и «е» вместо «ё» значения не имеют. Нераспознанная пометка "
    "не теряется: она попадает в отчёт проверки.",
    "11. «Тип грунта» и «Тип изоляции» выбираются из списка; значение действует от своей "
    "дистанции до следующей заполненной.",
)


@dataclass(frozen=True, slots=True)
class ListEntry:
    """Запись выпадающего списка шифров."""

    text: str  # что писать в «Комментарии»
    type_key: str
    role: str  # POINT / START / END


def list_entries(group: str | None = None) -> tuple[ListEntry, ...]:
    """Записи справочника по OBJECT_TYPES (при group - только одной группы)."""
    result: list[ListEntry] = []
    for t in OBJECT_TYPES:
        if group is not None and t.group != group:
            continue
        for pattern in t.patterns:
            base = FULL_FORMS.get(pattern, pattern)
            result.append(ListEntry(base, t.key, POINT))
            if t.extended:
                result.append(ListEntry(base + ROLE_SUFFIX[START], t.key, START))
                result.append(ListEntry(base + ROLE_SUFFIX[END], t.key, END))
    return tuple(result)


def _type_by_key() -> dict[str, ObjectType]:
    return {t.key: t for t in OBJECT_TYPES}


def _icon_image(type_key: str, max_h: int = 48, max_w: int = 80) -> XlImage | None:
    """Уменьшенный значок типа или None, если файла нет."""
    try:
        path = assets.icon_png(type_key)
    except KeyError:
        return None
    if not path.is_file():
        return None
    with Image.open(path) as src:
        img = src.convert("RGBA")
    scale = min(max_h / img.height, max_w / img.width, 1.0)
    size = (max(1, round(img.width * scale)), max(1, round(img.height * scale)))
    img = img.resize(size, Image.LANCZOS)
    buf = BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    xl = XlImage(buf)
    xl.width, xl.height = size
    return xl


_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
_WRAP = Alignment(wrap_text=True, vertical="center")


def _style_header(ws, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _fill_data_sheet(ws, lists_len: tuple[int, int, int]) -> None:
    for c, key in enumerate(TEMPLATE_KEYS, start=1):
        title = COLUMNS[key][0]
        ws.cell(row=1, column=c, value=title)
        ws.column_dimensions[get_column_letter(c)].width = 34 if key == "comment" else 16
    _style_header(ws, 1, len(TEMPLATE_KEYS))
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"

    def col(key: str) -> str:
        return get_column_letter(TEMPLATE_KEYS.index(key) + 1)

    n_objects, n_soil, n_ins = lists_len

    comment = DataValidation(
        type="list",
        formula1=f"'{SHEET_LISTS}'!$A$1:$A${n_objects}",
        allow_blank=True,
        showErrorMessage=False,
        showInputMessage=True,
        promptTitle=COMMENT_PROMPT_TITLE,
        prompt=COMMENT_PROMPT,
    )
    soil = DataValidation(
        type="list",
        formula1=f"'{SHEET_LISTS}'!$B$1:$B${n_soil}",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Тип грунта",
        error="Выберите тип грунта из списка: " + ", ".join(SOIL_TYPES) + ".",
    )
    ins = DataValidation(
        type="list",
        formula1=f"'{SHEET_LISTS}'!$C$1:$C${n_ins}",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Тип изоляции",
        error="Выберите тип изоляции из списка: " + ", ".join(INSULATION_TYPES) + ".",
    )
    for dv, key in ((comment, "comment"), (soil, "soil_type"), (ins, "ins_type")):
        letter = col(key)
        dv.add(f"{letter}2:{letter}{LAST_ROW}")
        ws.add_data_validation(dv)


def _fill_objects_sheet(ws) -> None:
    headers = ("Что писать в Комментарии", "Объект", "Вид на схеме", "Значок")
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(headers))
    for letter, width in zip("ABCD", (34, 36, 26, 14), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    types = _type_by_key()
    row = 2
    blocks = ((LEGEND, "Объекты новой легенды"), (LEGACY, "Объекты старого словаря"))
    for group, title in blocks:
        ws.cell(row=row, column=1, value=title).font = _HEADER_FONT
        for c in range(1, 5):
            ws.cell(row=row, column=c).fill = _HEADER_FILL
        row += 1
        images: dict[str, XlImage | None] = {}
        for entry in list_entries(group):
            t = types[entry.type_key]
            ws.cell(row=row, column=1, value=entry.text)
            ws.cell(row=row, column=2, value=t.title)
            ws.cell(row=row, column=3, value=PLACEMENT_TITLES[t.placement])
            for c in range(1, 4):
                ws.cell(row=row, column=c).alignment = _WRAP
            # Одна картинка на тип: на первой его строке, остальные строки типа пустые.
            first = entry.type_key not in images
            if first:
                img = _icon_image(entry.type_key)
                images[entry.type_key] = img
                if img is not None:
                    ws.add_image(img, f"D{row}")
            shown = images[entry.type_key] if first else None
            ws.row_dimensions[row].height = max(18.0, shown.height * 0.75 + 4) if shown else 18.0
            row += 1


def _fill_howto_sheet(ws) -> None:
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(HOWTO_LINES, start=1):
        cell = ws.cell(row=i, column=1, value=line or None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=13)


def _fill_lists_sheet(ws) -> tuple[int, int, int]:
    entries = [e.text for e in list_entries()]
    for r, text in enumerate(entries, start=1):
        ws.cell(row=r, column=1, value=text)
    for r, text in enumerate(SOIL_TYPES, start=1):
        ws.cell(row=r, column=2, value=text)
    for r, text in enumerate(INSULATION_TYPES, start=1):
        ws.cell(row=r, column=3, value=text)
    return len(entries), len(SOIL_TYPES), len(INSULATION_TYPES)


def build_template(path: Path, sheet_name: str = "Result") -> None:
    """Собирает шаблон книги и сохраняет в `path` (листы: данные, объекты, правила, списки)."""
    wb = Workbook()
    data = wb.active
    data.title = sheet_name
    objects = wb.create_sheet(SHEET_OBJECTS)
    howto = wb.create_sheet(SHEET_HOWTO)
    lists = wb.create_sheet(SHEET_LISTS)

    lens = _fill_lists_sheet(lists)
    lists.sheet_state = "hidden"
    _fill_data_sheet(data, lens)
    _fill_objects_sheet(objects)
    _fill_howto_sheet(howto)
    wb.active = 0

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


__all__ = ["TEMPLATE_KEYS", "ListEntry", "build_template", "list_entries"]
