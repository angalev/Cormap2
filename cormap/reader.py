"""Чтение книги Excel (шаблон версии 1 и 2) и проверка данных.

Колонки ищутся по тексту заголовка (строка 1), а не по букве.
"""

from __future__ import annotations

import math
import zipfile
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from cormap.model import (
    INSULATION_TYPES,
    SOIL_TYPES,
    BlockSummary,
    CheckReport,
    CommentBlock,
    CommentItem,
    DefectsBlock,
    DepthBlock,
    GpsBlock,
    GpsItem,
    GradientBlock,
    InsulationBlock,
    IntensiveBlock,
    SoilBlock,
    Survey,
    TypeMark,
    normalize_text,
)

MAX_FILE_BYTES = 50 * 1024 * 1024  # 50 МБ (ТЗ 0.4, раздел 9)
MAX_DATA_ROWS = 400_000  # трасса 300 км с шагом 1 м с запасом
MAX_ROUTE_M = 300_000.0  # длиннее - предупреждение route_too_long (ТЗ, раздел 4.1)
MAX_UNPACKED_BYTES = 1024 * 1024 * 1024  # защита от «бомбы» в zip

# Лист данных (ТЗ 4.1): `Result` - после модуля refinery, `Template` - без обработки.
SHEET_NAME = "Result"
SHEET_TEMPLATE = "Template"
# Книга, не прошедшая через модуль: читается `Template`, а если его нет - `Result`.
RAW_SHEETS: tuple[str, ...] = (SHEET_TEMPLATE, SHEET_NAME)

# Ключ -> (заголовок как в ТЗ, обязательная ли колонка)
COLUMNS: dict[str, tuple[str, bool]] = {
    "dist": ("Дистанция, м", True),
    "u_on": ("U т-з, В", True),
    "u_off": ("U т-з выкл, В", True),
    "gl_on": ("Гр-нт лев. вкл, мВ", False),
    "gl_off": ("Гр-нт лев. выкл, мВ", False),
    "gr_on": ("Гр-нт прав. вкл, мВ", False),
    "gr_off": ("Гр-нт прав. выкл, мВ", False),
    "comment": ("Комментарий", False),
    "gps": ("GPS", False),
    "soil_dist": ("Дистанция грун, м", False),
    "rho": ("Уд. соп-е грунтов, Ом*м", False),
    "soil_type": ("Тип грунта", False),
    "depth_dist": ("Дистанция глубины, м", False),
    "depth": ("Глубина заложения трубы, м", False),
    "dmg_dist": ("Дистанция повреждений, м", False),
    "dmg": ("Повреждения ИП", False),
    "ins_dist": ("Дистанция изоляции, м", False),
    "ins_type": ("Тип изоляции", False),
}
_BY_HEADER = {normalize_text(title): key for key, (title, _) in COLUMNS.items()}
REQUIRED = tuple(key for key, (_, req) in COLUMNS.items() if req)

# Блоки с собственной колонкой дистанции: (название, ключ дистанции, ключи значений)
_OWN_DIST_BLOCKS = (
    ("Грунт", "soil_dist", ("rho", "soil_type")),
    ("Глубина", "depth_dist", ("depth",)),
    ("Повреждения", "dmg_dist", ("dmg",)),
    ("Изоляция", "ins_dist", ("ins_type",)),
)

_EMPTY, _OK, _BAD = 0, 1, 2


def _parse_number(value: Any) -> tuple[int, float]:
    """Число из ячейки: число или строка с десятичной запятой."""
    if value is None:
        return _EMPTY, 0.0
    if isinstance(value, bool):
        return _BAD, 0.0
    if isinstance(value, int | float):
        return (_OK, float(value)) if math.isfinite(value) else (_BAD, 0.0)
    if isinstance(value, str):
        text = value.replace(" ", " ").strip().replace(" ", "")
        if not text:
            return _EMPTY, 0.0
        try:
            number = float(text.replace(",", "."))
        except ValueError:
            return _BAD, 0.0
        return (_OK, number) if math.isfinite(number) else (_BAD, 0.0)
    return _BAD, 0.0


def _text(value: Any) -> str | None:
    """Непустой текст ячейки (как есть) или None."""
    if value is None:
        return None
    if isinstance(value, str):
        return value if value.strip() else None
    return str(value)


def _row_is_blank(row: tuple[Any, ...]) -> bool:
    for v in row:
        if v is not None and (not isinstance(v, str) or v.strip()):
            return False
    return True


def _header_map(header_row: tuple[Any, ...]) -> tuple[dict[str, int], list[tuple[int, str]]]:
    """Ключ колонки -> индекс (первое вхождение); список (индекс, текст) повторов."""
    found: dict[str, int] = {}
    duplicates: list[tuple[int, str]] = []
    for idx, value in enumerate(header_row):
        if value is None or not str(value).strip():
            continue
        key = _BY_HEADER.get(normalize_text(value))
        if key is None:
            continue
        if key in found:
            duplicates.append((idx, str(value)))
        else:
            found[key] = idx
    return found, duplicates


def _missing(found: dict[str, int]) -> list[str]:
    return [COLUMNS[k][0] for k in REQUIRED if k not in found]


def _fail(report: CheckReport, code: str, message: str) -> tuple[None, CheckReport]:
    report.error(code, message)
    return None, report


def _check_file(path: Path, max_bytes: int, report: CheckReport) -> bool:
    if path.suffix.lower() != ".xlsx":
        report.error("bad_extension", "Принимаются только файлы формата .xlsx.")
        return False
    if not path.is_file():
        report.error("not_found", f"Файл не найден: {path}")
        return False
    size = path.stat().st_size
    if size > max_bytes:
        report.error(
            "file_too_big",
            f"Файл слишком большой: {size / 1024 / 1024:.1f} МБ, допустимо не более "
            f"{max_bytes / 1024 / 1024:.0f} МБ.",
        )
        return False
    try:
        with zipfile.ZipFile(path) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_UNPACKED_BYTES:
                report.error("file_too_big", "Содержимое файла после распаковки слишком велико.")
                return False
    except zipfile.BadZipFile:
        report.error("bad_file", "Файл повреждён или не является книгой .xlsx.")
        return False
    return True


def file_problem(path: str | Path, max_bytes: int = MAX_FILE_BYTES) -> tuple[str, str] | None:
    """Почему файл нельзя открывать: (код, текст) первой причины; None - можно.

    Те же пределы, что при чтении книги (расширение, размер, размер после распаковки).
    Нужна там, где книгу открывает не `read_workbook` (модуль refinery).
    """
    report = CheckReport(path=str(path))
    if _check_file(Path(path), max_bytes, report):
        return None
    issue = report.errors[0]
    return issue.code, issue.message


def read_workbook(
    path: str | Path,
    *,
    sheets: Sequence[str] = (SHEET_NAME,),
    max_rows: int = MAX_DATA_ROWS,
    max_bytes: int = MAX_FILE_BYTES,
) -> tuple[Survey | None, CheckReport]:
    """Читает книгу и проверяет её.

    ``sheets`` - имена листа данных по порядку предпочтения: читается первый найденный.

    Возвращает (Survey, отчёт). Survey равен None, если книгу прочитать нельзя
    (не .xlsx, лимиты, нет листа данных или его обязательных колонок). При ошибках в данных
    Survey возвращается, но `report.ok` ложно - строить карту нельзя.
    """
    path = Path(path)
    report = CheckReport(path=str(path))
    if not _check_file(path, max_bytes, report):
        return None, report
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl бросает разные исключения на битых файлах
        return _fail(report, "bad_file", f"Не удалось открыть книгу: {exc}")
    try:
        return _read(wb, report, max_rows, tuple(sheets))
    finally:
        wb.close()


def _first_row(ws: Any) -> tuple[Any, ...]:
    ws.reset_dimensions()
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        return row
    return ()


def _choose_sheet(
    wb: Any, report: CheckReport, sheets: tuple[str, ...]
) -> tuple[Any, dict[str, int], list[tuple[int, str]], tuple[Any, ...]] | None:
    """Первый найденный лист из ``sheets`` (без учёта регистра и крайних пробелов).

    Остальные листы книги не читаются.
    """
    by_title = {ws.title.strip().casefold(): ws for ws in reversed(wb.worksheets)}
    ws = next((by_title[n.casefold()] for n in sheets if n.casefold() in by_title), None)
    if ws is None:
        names = ", ".join(wb.sheetnames)
        wanted = " или ".join(f"«{n}»" for n in sheets)
        report.error("no_sheet", f"В книге нет листа {wanted}. Листы книги: {names}.")
        return None
    header = _first_row(ws)
    found, dups = _header_map(header)
    missing = _missing(found)
    if missing:
        cols = ", ".join(f"«{m}»" for m in missing)
        report.error("missing_columns", f"Нет обязательных колонок на листе «{ws.title}»: {cols}.")
        return None
    return ws, found, dups, header


def _read(
    wb: Any, report: CheckReport, max_rows: int, sheets: tuple[str, ...]
) -> tuple[Survey | None, CheckReport]:
    chosen = _choose_sheet(wb, report, sheets)
    if chosen is None:
        return None, report
    ws, cols, duplicates, header = chosen
    report.sheet = ws.title

    for idx, text in duplicates:
        report.warning(
            "duplicate_column",
            f"Колонка «{text}» повторяется, используется первая.",
            cell=f"{get_column_letter(idx + 1)}1",
            row=1,
        )
    known_idx = set(cols.values()) | {i for i, _ in duplicates}
    for idx, value in enumerate(header):
        if idx not in known_idx and value is not None and str(value).strip():
            report.warning(
                "unknown_column",
                f"Колонка «{value}» не распознана и не читается.",
                cell=f"{get_column_letter(idx + 1)}1",
                row=1,
            )
    for title, dist_key, value_keys in _OWN_DIST_BLOCKS:
        have_values = [k for k in value_keys if k in cols]
        if dist_key not in cols and have_values:
            report.warning(
                "block_no_distance",
                f"Блок «{title}»: нет колонки «{COLUMNS[dist_key][0]}», данные блока не читаются.",
            )
        elif dist_key in cols and not have_values:
            report.warning(
                "block_no_values",
                f"Блок «{title}»: есть колонка «{COLUMNS[dist_key][0]}», "
                "но нет колонок со значениями.",
            )

    survey = _read_rows(ws, cols, report, max_rows)
    return survey, report


def _read_rows(ws: Any, cols: dict[str, int], report: CheckReport, max_rows: int) -> Survey | None:
    idx = {k: cols.get(k, -1) for k in COLUMNS}
    letters = {k: get_column_letter(i + 1) for k, i in cols.items()}

    def cell_of(key: str, rownum: int) -> str:
        return f"{letters[key]}{rownum}"

    def title(key: str) -> str:
        return COLUMNS[key][0]

    def get(row: tuple[Any, ...], key: str) -> Any:
        i = idx[key]
        return row[i] if 0 <= i < len(row) else None

    # Последняя принятая дистанция по каждой колонке дистанции (проверка убывания).
    last_dist: dict[str, float] = {}

    def num(row: tuple[Any, ...], key: str, rownum: int, error: bool) -> float | None:
        status, value = _parse_number(get(row, key))
        if status == _OK:
            return value
        if status == _BAD:
            raw = get(row, key)
            msg = f"в колонке «{title(key)}» нечисловое значение «{raw}»."
            if error:
                report.error("not_a_number", msg[0].upper() + msg[1:], cell_of(key, rownum), rownum)
            else:
                report.warning(
                    "not_a_number",
                    msg[0].upper() + msg[1:] + " Значение пропущено.",
                    cell_of(key, rownum),
                    rownum,
                )
        return None

    def dist_of(row: tuple[Any, ...], key: str, rownum: int, error: bool) -> float | None:
        """Дистанция строки с проверкой на убывание внутри колонки."""
        d = num(row, key, rownum, error)
        if d is None:
            return None
        prev = last_dist.get(key)
        if prev is not None and d < prev:
            msg = f"Дистанция убывает: {d:g} после {prev:g} (колонка «{title(key)}»)."
            if error:
                report.error("distance_decreases", msg, cell_of(key, rownum), rownum)
            else:
                report.warning("distance_decreases", msg, cell_of(key, rownum), rownum)
        last_dist[key] = d
        return d

    main_dist: list[float] = []
    u_on: list[float | None] = []
    u_off: list[float | None] = []
    g_dist: list[float] = []
    g_vals: tuple[list[float | None], ...] = ([], [], [], [])
    comments: list[CommentItem] = []
    gps: list[GpsItem] = []
    soil_dist: list[float] = []
    soil_rho: list[float] = []
    soil_types: list[TypeMark] = []
    soil_rows: list[float] = []
    depth_dist: list[float] = []
    depth_val: list[float] = []
    dmg_dist: list[float] = []
    dmg_val: list[float | str] = []
    ins_types: list[TypeMark] = []
    comment_rows: list[float] = []
    gps_rows: list[float] = []

    def type_mark(
        row: tuple[Any, ...],
        key: str,
        dist: float,
        rownum: int,
        allowed: tuple[str, ...],
        what: str,
    ) -> TypeMark | None:
        raw = _text(get(row, key))
        if raw is None:
            return None
        value = normalize_text(raw)
        known = value in allowed
        if not known:
            report.warning(
                "unknown_type",
                f"Неизвестный {what} «{raw.strip()}». Допустимые значения: {', '.join(allowed)}. "
                "Участок будет нарисован серой штриховкой.",
                cell_of(key, rownum),
                rownum,
            )
        return TypeMark(dist, value, known, rownum)

    grad_keys = ("gl_on", "gl_off", "gr_on", "gr_off")
    has_grad = any(k in cols for k in grad_keys)
    main_keys = ("dist", "u_on", "u_off")
    soil_vals = [k for k in ("rho", "soil_type") if k in cols]
    ins_vals = "ins_type" in cols

    rownum = 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        rownum += 1
        if _row_is_blank(row):
            continue
        if rownum - 1 > max_rows:
            report.error(
                "too_many_rows",
                f"В книге больше {max_rows} строк данных: обработка прекращена.",
                row=rownum,
            )
            return None

        # Основная дистанция и всё, что к ней привязано.
        d = dist_of(row, "dist", rownum, error=True)
        uo = num(row, "u_on", rownum, error=True)
        uf = num(row, "u_off", rownum, error=True)
        grads = [num(row, k, rownum, error=False) if k in cols else None for k in grad_keys]
        comment = _text(get(row, "comment"))
        gps_text = _text(get(row, "gps"))
        if d is None:
            if _parse_number(get(row, "dist"))[0] == _EMPTY:
                if any(_parse_number(get(row, k))[0] != _EMPTY for k in main_keys[1:]):
                    report.error(
                        "empty_value",
                        f"Не заполнена «{title('dist')}» при заполненных измерениях.",
                        cell_of("dist", rownum),
                        rownum,
                    )
                elif (
                    comment is not None or gps_text is not None or any(g is not None for g in grads)
                ):
                    report.warning(
                        "no_distance",
                        "Нет дистанции: комментарий, GPS и градиенты строки пропущены.",
                        cell_of("dist", rownum),
                        rownum,
                    )
        else:
            if uo is not None or uf is not None:
                main_dist.append(d)
                u_on.append(uo)
                u_off.append(uf)
            if has_grad and any(g is not None for g in grads):
                g_dist.append(d)
                for lst, val in zip(g_vals, grads, strict=True):
                    lst.append(val)
            if comment is not None:
                comments.append(CommentItem(d, rownum, comment))
                comment_rows.append(d)
            if gps_text is not None:
                gps.append(GpsItem(d, rownum, gps_text))
                gps_rows.append(d)

        # Грунт.
        if "soil_dist" in cols and soil_vals:
            sd = dist_of(row, "soil_dist", rownum, error=False)
            rho = num(row, "rho", rownum, error=False) if "rho" in cols else None
            stype = _text(get(row, "soil_type")) if "soil_type" in cols else None
            if sd is not None:
                if rho is not None:
                    soil_dist.append(sd)
                    soil_rho.append(rho)
                if stype is not None:
                    mark = type_mark(row, "soil_type", sd, rownum, SOIL_TYPES, "тип грунта")
                    if mark:
                        soil_types.append(mark)
                if rho is not None or stype is not None:
                    soil_rows.append(sd)
            elif rho is not None or stype is not None:
                _no_dist(report, cell_of("soil_dist", rownum), rownum, "Грунт")

        # Глубина.
        if "depth_dist" in cols and "depth" in cols:
            dd = dist_of(row, "depth_dist", rownum, error=False)
            dv = num(row, "depth", rownum, error=False)
            if dd is not None and dv is not None:
                depth_dist.append(dd)
                depth_val.append(dv)
            elif dd is None and dv is not None:
                _no_dist(report, cell_of("depth_dist", rownum), rownum, "Глубина")

        # Повреждения (ручные): значение хранится как есть.
        if "dmg_dist" in cols and "dmg" in cols:
            md = dist_of(row, "dmg_dist", rownum, error=False)
            raw = get(row, "dmg")
            status, dnum = _parse_number(raw)
            dval: float | str | None = (
                dnum if status == _OK else _text(raw) if status == _BAD else None
            )
            if md is not None and dval is not None:
                dmg_dist.append(md)
                dmg_val.append(dval)
            elif md is None and dval is not None:
                _no_dist(report, cell_of("dmg_dist", rownum), rownum, "Повреждения")

        # Изоляция.
        if "ins_dist" in cols and ins_vals:
            idist = dist_of(row, "ins_dist", rownum, error=False)
            itext = _text(get(row, "ins_type"))
            if idist is not None and itext is not None:
                mark = type_mark(row, "ins_type", idist, rownum, INSULATION_TYPES, "тип изоляции")
                if mark:
                    ins_types.append(mark)
            elif idist is None and itext is not None:
                _no_dist(report, cell_of("ins_dist", rownum), rownum, "Изоляция")

    survey = Survey(
        sheet_name=ws.title,
        intensive=IntensiveBlock(tuple(main_dist), tuple(u_on), tuple(u_off)),
        gradients=GradientBlock(tuple(g_dist), *(tuple(v) for v in g_vals)),
        comments=CommentBlock(tuple(comments)),
        gps=GpsBlock(tuple(gps)),
        soil=SoilBlock(tuple(soil_dist), tuple(soil_rho), tuple(soil_types)),
        depth=DepthBlock(tuple(depth_dist), tuple(depth_val)),
        defects=DefectsBlock(tuple(dmg_dist), tuple(dmg_val)),
        insulation=InsulationBlock(tuple(ins_types)),
    )
    if survey.distance_range is not None:
        lo, hi = survey.distance_range
        if hi - lo > MAX_ROUTE_M:
            report.warning(
                "route_too_long",
                f"Трасса длиннее 300 км (диапазон дистанций {(hi - lo) / 1000:.1f} км): "
                "построение возможно, но займёт много времени и листов.",
            )
    if not main_dist:
        report.error(
            "no_data", "В блоке «Интенсивные измерения» нет данных: строить карту не из чего."
        )

    n_on = sum(v is not None for v in u_on)
    n_off = sum(v is not None for v in u_off)
    if main_dist:
        for key, count in (("u_on", n_on), ("u_off", n_off)):
            if count == 0:
                report.warning(
                    "required_column_empty",
                    f"Колонка «{title(key)}» есть, но в ней нет ни одного значения.",
                    cell=f"{letters[key]}1",
                    row=1,
                )
    report.blocks = [
        BlockSummary.from_distances(
            "Интенсивные измерения", main_dist, f"U т-з: {n_on}, U т-з выкл: {n_off} значений"
        ),
        BlockSummary.from_distances("Градиенты", g_dist),
        BlockSummary.from_distances("Комментарии", comment_rows),
        BlockSummary.from_distances("GPS", gps_rows),
        BlockSummary.from_distances(
            "Грунт",
            soil_rows,
            f"УЭС: {len(soil_dist)}, типов грунта: {len(soil_types)}",
        ),
        BlockSummary.from_distances("Глубина", depth_dist),
        BlockSummary.from_distances("Повреждения", dmg_dist),
        BlockSummary.from_distances("Изоляция", [t.dist for t in ins_types]),
    ]
    return survey


def _no_dist(report: CheckReport, cell: str, rownum: int, block: str) -> None:
    report.warning(
        "no_distance",
        f"Блок «{block}»: значение без дистанции пропущено.",
        cell,
        rownum,
    )


__all__ = [
    "COLUMNS",
    "RAW_SHEETS",
    "REQUIRED",
    "SHEET_NAME",
    "SHEET_TEMPLATE",
    "read_workbook",
]
