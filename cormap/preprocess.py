"""Предварительная обработка книги модулем заказчика `refinery` (ТЗ, раздел 4.5).

Модуль - файл `cormap/refinery.py`, который кладёт заказчик: сервис знает только его
интерфейс и импортирует его в момент вызова. Здесь - настройки шага, проверка правил
до вызова и сам вызов с разбором всех исходов. Django не используется.
"""

from __future__ import annotations

import importlib
import logging
from dataclasses import dataclass, fields
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from cormap.reader import MAX_DATA_ROWS, file_problem

logger = logging.getLogger("cormap.preprocess")

SOURCE_SHEET = "Template"  # лист, который читает модуль (ТЗ 4.5), если он не назвал свой

# Флажки по блокам: (заголовок блока, ((переменная, подпись), ...)). Подписи - из ТЗ 4.5.
FLAG_GROUPS: tuple[tuple[str, tuple[tuple[str, str], ...]], ...] = (
    (
        "U т-з",
        (
            ("utz", "U т-з вкл"),
            ("utzd", "U т-з выкл"),
            ("utzdc", "U т-з выкл (генератор)"),
        ),
    ),
    (
        "Левый градиент",
        (
            ("grlu", "Левый градиент вкл (интерполяция)"),
            ("grld", "Левый градиент выкл (интерполяция)"),
            ("grluc", "Левый градиент вкл (генератор)"),
            ("grldc", "Левый градиент выкл (0.8 от градиента вкл)"),
        ),
    ),
    (
        "Правый градиент",
        (
            ("grru", "Правый градиент вкл (интерполяция)"),
            ("grrd", "Правый градиент выкл (интерполяция)"),
            ("grruc", "Правый градиент вкл (генератор)"),
            ("grrdc", "Правый градиент выкл (0.8 от градиента вкл)"),
        ),
    ),
)

FLAG_NAMES: tuple[str, ...] = tuple(name for _, items in FLAG_GROUPS for name, _ in items)
FLAG_LABELS: dict[str, str] = {name: label for _, items in FLAG_GROUPS for name, label in items}

STEP_MIN = 1
STEP_MAX = 10
STEP_DEFAULT = 5
STEP_LABEL = "Шаг интерполяции, м"


@dataclass(frozen=True)
class RefineOptions:
    """Настройки шага обработки: 11 флажков, шаг интерполяции и сид."""

    utz: bool = False
    utzd: bool = False
    utzdc: bool = False
    grlu: bool = False
    grld: bool = False
    grluc: bool = False
    grldc: bool = False
    grru: bool = False
    grrd: bool = False
    grruc: bool = False
    grrdc: bool = False
    step: int = STEP_DEFAULT
    random_seed: int | None = None

    def flags(self) -> dict[str, bool]:
        """Только флажки (имя -> значение)."""
        return {name: bool(getattr(self, name)) for name in FLAG_NAMES}

    @property
    def any_flag(self) -> bool:
        """Выставлен хотя бы один флажок: без флажков модуль не вызывается (ТЗ, 4.5)."""
        return any(self.flags().values())

    def enabled_labels(self) -> list[str]:
        """Подписи выставленных флажков в порядке таблицы ТЗ."""
        return [FLAG_LABELS[name] for name in FLAG_NAMES if getattr(self, name)]

    def to_dict(self) -> dict[str, Any]:
        """Словарь для хранения в JSON (параметры карты)."""
        return {f.name: getattr(self, f.name) for f in fields(self)}

    @classmethod
    def from_dict(cls, data: dict[str, Any] | None) -> RefineOptions:
        """Обратное преобразование; лишние ключи отбрасываются, недостающие - по умолчанию."""
        data = data or {}
        known = {f.name for f in fields(cls)}
        return cls(**{k: v for k, v in data.items() if k in known})


def validate(options: RefineOptions) -> list[tuple[str, str]]:
    """Нарушения правил ТЗ 4.5 парами (имя поля, текст). Пустое имя - общая ошибка.

    Ни одного флажка - не нарушение: тогда модуль просто не вызывается.
    """
    o = options
    errors: list[tuple[str, str]] = []
    if o.utzdc and not o.utz:
        errors.append(("utzdc", f"«{FLAG_LABELS['utzdc']}» возможно только вместе с «U т-з вкл»."))
    if o.utzdc and o.utzd:
        errors.append(("utzdc", f"«{FLAG_LABELS['utzdc']}» нельзя вместе с «U т-з выкл»."))
    for on, off, on_c, off_c in (
        ("grlu", "grld", "grluc", "grldc"),
        ("grru", "grrd", "grruc", "grrdc"),
    ):
        if getattr(o, on) and getattr(o, on_c):
            errors.append((on_c, f"«{FLAG_LABELS[on]}» и «{FLAG_LABELS[on_c]}» вместе нельзя."))
        if getattr(o, off) and getattr(o, off_c):
            errors.append((off_c, f"«{FLAG_LABELS[off]}» и «{FLAG_LABELS[off_c]}» вместе нельзя."))
        if getattr(o, off_c) and not (getattr(o, on) or getattr(o, on_c)):
            errors.append(
                (
                    off_c,
                    f"«{FLAG_LABELS[off_c]}» возможно только при «{FLAG_LABELS[on]}» "
                    f"или «{FLAG_LABELS[on_c]}».",
                )
            )
    step = o.step
    if isinstance(step, bool) or not isinstance(step, int) or not STEP_MIN <= step <= STEP_MAX:
        errors.append(("step", f"Шаг интерполяции - целое число от {STEP_MIN} до {STEP_MAX}."))
    return errors


@dataclass(frozen=True)
class RefineStats:
    """Статистика вызова модуля (всё сериализуется в JSON)."""

    input_rows: Any = None
    output_rows: Any = None
    distance_range_m: Any = None
    columns_generated: Any = None
    calculation_time_ms: Any = None

    def to_dict(self) -> dict[str, Any]:
        return {f.name: _jsonable(getattr(self, f.name)) for f in fields(self)}


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, bool | int | float | str):
        return value
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, list | tuple | set | frozenset):
        return [_jsonable(v) for v in value]
    return str(value)


@dataclass(frozen=True)
class RefineOutcome:
    """Итог вызова модуля."""

    ok: bool
    error_code: str = ""
    error_message: str = ""
    stats: RefineStats | None = None


# Где искать модуль: отдельно установленный `refinery` важнее файла в ядре.
MODULE_NAMES: tuple[str, ...] = ("refinery", "cormap.refinery")


def _import_refinery():
    """Ленивый импорт модуля заказчика; ImportError, если его нет ни под одним именем."""
    for name in MODULE_NAMES[:-1]:
        try:
            return importlib.import_module(name)
        except ImportError:
            continue
    return importlib.import_module(MODULE_NAMES[-1])


def refinery_available() -> bool:
    """Модуль `refinery` можно импортировать."""
    try:
        _import_refinery()
    except ImportError:
        return False
    return True


def _fail(code: str, message: str) -> RefineOutcome:
    return RefineOutcome(ok=False, error_code=code, error_message=message)


def run_refinery_step(src: Path, dst: Path, options: RefineOptions) -> RefineOutcome:
    """Прогоняет книгу `src` через модуль, результат пишет в `dst`. Исключений не бросает."""
    try:
        module = _import_refinery()
    except ImportError:
        return _fail("not_installed", "Модуль предварительной обработки (refinery) не установлен.")
    src, dst = Path(src), Path(dst)
    # модуль загружает книгу в память целиком: пределы размера проверяем до вызова
    problem = file_problem(src)
    if problem is not None:
        return _fail(*problem)
    trimmed = dst.with_name(dst.stem + "_trimmed.xlsx")
    try:
        dst.parent.mkdir(parents=True, exist_ok=True)
        sheet = str(getattr(module, "SOURCE_SHEET", SOURCE_SHEET))
        try:
            source = _without_empty_tail(src, trimmed, sheet)
        except _TooManyRows:
            return _fail(
                "too_many_rows",
                f"В книге больше {MAX_DATA_ROWS} строк данных: обработка невозможна.",
            )
        settings = module.RefinerySettings(
            **options.flags(), step=options.step, random_seed=options.random_seed
        )
        result = module.run_refinery(input_file=source, settings=settings, output_file=dst)
    except Exception as exc:  # модуль не должен бросать, но сервис это не предполагает
        logger.exception("Исключение при вызове refinery, книга %s", src.name)
        _discard(dst)
        text = str(exc).strip() or type(exc).__name__
        return _fail("calculation_error", f"Сбой при обработке: {text}")
    finally:
        _discard(trimmed)

    if not getattr(result, "success", False):
        _discard(dst)
        code = str(getattr(result, "error_code", "") or "unknown_error")
        message = str(getattr(result, "error_message", "") or "Модуль не сообщил причину ошибки.")
        return _fail(code, message)
    if not dst.is_file():
        return _fail("write_error", "Модуль сообщил об успехе, но выходной файл не создан.")

    raw = getattr(result, "stats", None)
    stats = None
    if raw is not None:
        stats = RefineStats(**{f.name: getattr(raw, f.name, None) for f in fields(RefineStats)})
    return RefineOutcome(ok=True, stats=stats)


class _TooManyRows(Exception):
    pass


def _without_empty_tail(src: Path, trimmed: Path, sheet: str) -> Path:
    """Книга для модуля без пустого «хвоста» листа данных.

    Excel нередко оставляет оформленную, но пустую ячейку далеко внизу листа (вплоть до
    последней строки, 1 048 576). Лист тогда числится огромным, и модуль, перебирая
    строки до последней, расходует гигабайты памяти на книге в десяток строк. Если лист
    объявлен длиннее допустимого, считаем настоящие строки: когда данных в пределах
    нормы, модулю отдаётся копия книги без хвоста (`trimmed`, значения вместо формул),
    иначе - _TooManyRows. Обычная книга возвращается как есть.
    """
    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return src  # об отсутствии листа сообщит модуль
        ws = wb[sheet]
        declared = ws.max_row
        if declared is not None and declared <= MAX_DATA_ROWS + 1:
            return src
        last = 0
        for number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if any(value is not None for value in row):
                last = number
                if last > MAX_DATA_ROWS + 1:
                    raise _TooManyRows
    finally:
        wb.close()
    if declared is not None and last >= declared:
        return src
    full = load_workbook(src, data_only=True)
    try:
        ws = full[sheet]
        if ws.max_row <= last:
            return src
        ws.delete_rows(last + 1, ws.max_row - last)
        full.save(trimmed)
    finally:
        full.close()
    logger.info(
        "Книга %s: лист «%s» объявлен на %s строк, данных %d - пустой хвост отброшен",
        src.name,
        sheet,
        declared,
        last,
    )
    return trimmed


def _discard(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass
