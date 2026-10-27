"""Тестовые книги для проверки сайта в браузере: python tools/make_demo_books.py.

Результат - каталог out/demo/ (не в репозитории).
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tests.synth import make_sample_like, make_workbook  # noqa: E402

OUT = ROOT / "out" / "demo"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    make_sample_like(OUT / "obrazec_5km.xlsx")

    # Предупреждения, но строится: нераспознанные пометки, непарное «начало»,
    # неизвестный тип грунта.
    path = make_workbook(
        OUT / "s_oshibkami.xlsx",
        length_m=10_000,
        step_m=2,
        notes=[
            (200, "КИП, КИП-1"),
            (1000, "странная пометка"),
            (1500, "что-то непонятное; кран"),
            (3000, "река начало, р. Упа"),  # без «конца»
            (5000, "овраг конец"),  # без «начала»
            (7000, "УКЗ, УКЗ-2"),
        ],
    )
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = headers.index("Тип грунта") + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, col).value:
            ws.cell(row, col).value = "глина"  # неизвестный тип грунта
            break
    wb.save(path)

    # Блокирующая: нет обязательной колонки «U т-з выкл, В».
    wb = Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append(["Дистанция, м", "U т-з, В", "Комментарий"])
    for i in range(50):
        ws.append([i * 2, -1.2, "кип" if i == 10 else None])
    wb.save(OUT / "blokiruyushchaya.xlsx")

    make_workbook(OUT / "trassa_50km.xlsx", length_m=50_000, step_m=2)
    print(f"Готово: {OUT}")


if __name__ == "__main__":
    main()
