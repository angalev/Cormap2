"""Сквозной сценарий: загрузка, проверка, правка разбиения, построение, файлы."""

from __future__ import annotations

import io
import json
import zipfile

from django.core.paginator import Paginator
from maps import storage, worker
from maps.models import AuditLog, MapJob
from openpyxl import load_workbook

from tests.synth import make_sample_like


def _content(response) -> str:
    return response.content.decode()


def test_full_scenario(client, users, login, small_book, post_new):
    login(users["a"])

    # 1. загрузка (километраж начинается с 12 км)
    response = post_new(small_book, start_km="12")
    job = MapJob.objects.get()
    assert response["Location"] == f"/maps/{job.pk}/"
    page = _content(client.get(response["Location"]))
    assert "Идёт проверка" in page
    assert "poll.js" in page and 'http-equiv="refresh"' in page  # опрос и запасное обновление

    # 2. обработчик проверяет книгу
    assert worker.process_all() == 1
    job.refresh_from_db()
    assert job.state == MapJob.State.CHECKED
    assert job.check_result["ok"] is True
    assert job.split["count"] == 1

    # 3. отчёт на странице
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert "Интенсивные измерения" in page
    assert ">601<" in page  # число точек: 6000 м / 10 м + 1
    assert "странная пометка" in page  # нераспознанная пометка видна
    assert "НЕ РАСПОЗНАНА" in page
    assert "Требуют внимания" in page
    assert "КИП-1" in page  # таблица «строка - пометка - как распознана»
    assert "Контрольно-измерительный пункт" in page
    assert "Пересчитать разбиение" in page and "Построить карту" in page
    assert "Разбиение на листы" in page and "12,00" in page  # километраж с начальным км

    # 4. пересчёт границами в км (по километражу, с учётом начального км)
    response = client.post(f"/maps/{job.pk}/split/", {"mode": "bounds", "value": "12; 15; 18"})
    assert response.status_code == 302
    job.refresh_from_db()
    assert job.params["split"] == [0.0, 3000.0, 6000.0]
    assert job.split["count"] == 2
    assert job.split["kind"] == "bounds"
    assert job.split["bounds_m"] == [0.0, 3000.0, 6000.0]
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert "12; 15; 18" in page  # поле показывает границы в км по километражу
    rows = job.split["table"]
    assert rows[0][1:3] == ["12,00", "15,00"] and rows[1][1:3] == ["15,00", "18,00"]

    # 5. убывающие границы: ошибка у поля, состояние не портится
    before = (job.params, job.split, job.state)
    response = client.post(f"/maps/{job.pk}/split/", {"mode": "bounds", "value": "12; 18; 15"})
    assert response.status_code == 400
    assert "должны возрастать" in _content(response)
    job.refresh_from_db()
    assert (job.params, job.split, job.state) == before

    # 6. построение
    assert client.post(f"/maps/{job.pk}/build/").status_code == 302
    job.refresh_from_db()
    assert job.state == MapJob.State.QUEUED
    assert "в очереди" in _content(client.get(f"/maps/{job.pk}/"))
    status = client.get(f"/maps/{job.pk}/status/").json()
    assert status["state"] == "queued" and status["in_progress"] is True

    assert worker.process_all() == 1
    job.refresh_from_db()
    assert job.state == MapJob.State.DONE, job.error_text
    assert job.sections_count == 2
    status = client.get(f"/maps/{job.pk}/status/").json()
    assert status == {
        "state": "done",
        "status": "готово",
        "in_progress": False,
        "stalled": False,
        "progress_done": job.progress_done,
        "progress_total": status["progress_total"],
        "error": "",
    }

    # 7. страница готовой карты и файлы
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert "готово" in page
    assert "Скачать PDF" in page and "Архив PNG" in page and "Excel с графиками" in page
    assert f"/maps/{job.pk}/preview/2/1/" in page
    assert "Перестроить с другим разбиением" in page

    pdf = client.get(f"/maps/{job.pk}/files/pdf/")
    pdf_bytes = b"".join(pdf.streaming_content)
    assert pdf_bytes.startswith(b"%PDF") and len(pdf_bytes) > 20_000

    archive = client.get(f"/maps/{job.pk}/files/png-zip/")
    with zipfile.ZipFile(io.BytesIO(b"".join(archive.streaming_content))) as zf:
        names = sorted(zf.namelist())
    assert names == ["list_01_1.png", "list_01_2.png", "list_02_1.png", "list_02_2.png"]

    xlsx = client.get(f"/maps/{job.pk}/files/xlsx/")
    wb = load_workbook(io.BytesIO(b"".join(xlsx.streaming_content)))
    assert "Result" in wb.sheetnames and "Карта" in wb.sheetnames

    thumb = client.get(f"/maps/{job.pk}/preview/1/1/")
    thumb_bytes = b"".join(thumb.streaming_content)
    full = client.get(f"/maps/{job.pk}/preview/1/1/?full=1")
    full_bytes = b"".join(full.streaming_content)
    assert thumb["Content-Type"] == "image/png" and thumb_bytes.startswith(b"\x89PNG")
    assert full_bytes.startswith(b"\x89PNG")
    assert len(full_bytes) > len(thumb_bytes) > 1000
    assert client.get(f"/maps/{job.pk}/preview/2/2/").status_code == 200
    assert client.get(f"/maps/{job.pk}/preview/3/1/").status_code == 404

    # 8. журнал
    actions = list(AuditLog.objects.filter(job=job).values_list("action", flat=True))
    for expected in ("upload", "check", "build", "download"):
        assert expected in actions, expected
    assert actions.count("download") >= 3
    assert AuditLog.objects.filter(action="build", details__contains="поставлено").exists()
    assert AuditLog.objects.filter(action="build", details__contains="построено").exists()

    # 9. «Перестроить с другим разбиением» возвращает на шаг проверки
    assert client.post(f"/maps/{job.pk}/rebuild/").status_code == 302
    job.refresh_from_db()
    assert job.state == MapJob.State.CHECKED
    assert client.get(f"/maps/{job.pk}/files/pdf/").status_code == 404  # не готово
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert "Правка разбиения" in page
    client.post(f"/maps/{job.pk}/split/", {"mode": "length", "value": "6"})
    job.refresh_from_db()
    assert job.params["split"] == 6000.0 and job.split["count"] == 1
    client.post(f"/maps/{job.pk}/build/")
    worker.process_all()
    job.refresh_from_db()
    assert job.state == MapJob.State.DONE and job.sections_count == 1
    assert not (storage.resolve(job.png_dir) / "list_02_1.png").exists()  # старое убрано
    assert client.get(f"/maps/{job.pk}/preview/2/1/").status_code == 404


def test_sample_book_builds(client, users, login, tmp_path, post_new):
    book = make_sample_like(tmp_path / "obrazec.xlsx")
    login(users["a"])
    post_new(book, object_name="Образец")
    worker.process_all()
    job = MapJob.objects.get()
    assert job.state == MapJob.State.CHECKED
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert ">2751<" in page  # 5500 м, шаг 2 м
    assert "Образец" in page
    client.post(f"/maps/{job.pk}/build/")
    worker.process_all()
    job.refresh_from_db()
    assert job.state == MapJob.State.DONE, job.error_text
    pdf = b"".join(client.get(f"/maps/{job.pk}/files/pdf/").streaming_content)
    assert pdf.startswith(b"%PDF") and len(pdf) > 50_000
    png = b"".join(client.get(f"/maps/{job.pk}/preview/1/2/").streaming_content)
    assert png.startswith(b"\x89PNG")


def test_blocking_error_book(client, users, login, blocking_book, post_new):
    login(users["a"])
    post_new(blocking_book)
    worker.process_all()
    job = MapJob.objects.get()
    assert job.state == MapJob.State.CHECK_FAILED
    assert "U т-з выкл" in job.error_text
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert "Построить карту" not in page
    assert "Пересчитать разбиение" not in page
    assert "Загрузить исправленную книгу" in page
    assert "Нет обязательных колонок" in page
    # построение вручную тоже невозможно
    client.post(f"/maps/{job.pk}/build/")
    job.refresh_from_db()
    assert job.state == MapJob.State.CHECK_FAILED


def test_replace_book_after_failure(client, users, login, blocking_book, small_book, post_new):
    login(users["a"])
    post_new(blocking_book)
    worker.process_all()
    job = MapJob.objects.get()
    old_path = storage.resolve(job.source_path)
    from conftest import upload_of

    response = client.post(f"/maps/{job.pk}/replace/", {"book": upload_of(small_book)})
    assert response.status_code == 302
    job.refresh_from_db()
    assert job.state == MapJob.State.UPLOADED
    assert job.original_name == "small.xlsx"
    assert not old_path.exists()
    worker.process_all()
    job.refresh_from_db()
    assert job.state == MapJob.State.CHECKED


def test_replace_updates_upload_time_not_creation_time(
    client, users, login, blocking_book, small_book, post_new
):
    """«Загружен» на странице карты - время нынешней книги; срок хранения идёт от создания."""
    from datetime import timedelta

    from conftest import upload_of
    from django.utils import timezone

    login(users["a"])
    post_new(blocking_book)
    worker.process_all()
    job = MapJob.objects.get()
    long_ago = timezone.now() - timedelta(days=3)
    MapJob.objects.filter(pk=job.pk).update(created_at=long_ago, uploaded_at=long_ago)
    client.post(f"/maps/{job.pk}/replace/", {"book": upload_of(small_book)})
    worker.process_all()
    job.refresh_from_db()
    assert job.created_at == long_ago
    assert timezone.now() - job.uploaded_at < timedelta(minutes=1)
    shown = timezone.localtime(job.uploaded_at).strftime("%d.%m.%Y %H:%M")
    old = timezone.localtime(long_ago).strftime("%d.%m.%Y")
    html = client.get(f"/maps/{job.pk}/").content.decode()
    assert f"загружен {shown}" in html and f"загружен {old}" not in html


def test_replace_rejects_bad_file(client, users, login, blocking_book, post_new):
    from django.core.files.uploadedfile import SimpleUploadedFile

    login(users["a"])
    post_new(blocking_book)
    worker.process_all()
    job = MapJob.objects.get()
    response = client.post(
        f"/maps/{job.pk}/replace/", {"book": SimpleUploadedFile("a.xlsx", b"text")}
    )
    assert response.status_code == 400
    job.refresh_from_db()
    assert job.state == MapJob.State.CHECK_FAILED


def test_split_modes_and_errors(client, users, login, make_job):
    job = make_job(users["a"], check=True)
    assert job.state == MapJob.State.CHECKED
    login(users["a"])
    url = f"/maps/{job.pk}/split/"
    bad = {
        ("bounds", "5"): "не меньше двух",
        ("bounds", "0; abc"): "Не число",
        ("length", ""): "Укажите значение",
        ("length", "1; 2"): "одна положительная длина",
        ("length", "-1"): "положительная",
        ("bounds", "0; 3; 3"): "возрастать",
    }
    for (mode, value), text in bad.items():
        response = client.post(url, {"mode": mode, "value": value})
        assert response.status_code == 400, (mode, value)
        assert text in _content(response), (mode, value)
    assert client.post(url, {"mode": "zzz", "value": ""}).status_code == 400
    job.refresh_from_db()
    assert job.state == MapJob.State.CHECKED and job.params["split"] == "auto"
    # режим «авто» возвращает автоматическое разбиение
    assert client.post(url, {"mode": "length", "value": "2"}).status_code == 302
    assert client.post(url, {"mode": "auto", "value": ""}).status_code == 302
    job.refresh_from_db()
    assert job.params["split"] == "auto" and job.split["kind"] == "auto"
    # слишком много участков: ошибка ядра у поля
    response = client.post(url, {"mode": "length", "value": "0,001"})
    assert response.status_code == 400
    assert "предел" in _content(response)


def test_split_not_available_before_check(client, users, login, make_job):
    job = make_job(users["a"])  # uploaded
    login(users["a"])
    client.post(f"/maps/{job.pk}/split/", {"mode": "length", "value": "2"})
    client.post(f"/maps/{job.pk}/build/")
    job.refresh_from_db()
    assert job.state == MapJob.State.UPLOADED


def test_double_build_click_queues_once(client, users, login, make_job):
    job = make_job(users["a"], check=True)
    login(users["a"])
    client.post(f"/maps/{job.pk}/build/")
    client.post(f"/maps/{job.pk}/build/")
    assert AuditLog.objects.filter(action="build", details__contains="поставлено").count() == 1


def test_records_pagination_and_unrecognized_on_top(client, users, login, make_job):
    job = make_job(users["a"], state=MapJob.State.CHECKED)
    records = [
        {"row": i + 2, "dist": float(i), "text": f"КИП-{i}", "recognized": True, "what": "КИП"}
        for i in range(250)
    ]
    records.append(
        {
            "row": 999,
            "dist": 999.0,
            "text": "загадка-999",
            "recognized": False,
            "what": "НЕ РАСПОЗНАНА",
        }
    )
    MapJob.objects.filter(pk=job.pk).update(
        check_result={
            "ok": True,
            "blocks": [],
            "issues": [],
            "suppressed": {},
            "notes": {
                "records": records,
                "recognized": 250,
                "unknown": 1,
                "unpaired": 1,
                "unpaired_issues": [
                    {
                        "level": "warning",
                        "code": "note_start_without_end",
                        "message": "Начало без конца",
                        "cell": "H20",
                        "row": 20,
                    }
                ],
            },
        },
        split={"count": 0, "header": [], "table": [], "issues": [], "kind": "auto", "paper": "A4"},
    )
    login(users["a"])
    assert Paginator(records, 100).num_pages == 3
    for page_no in (1, 3):
        text = _content(client.get(f"/maps/{job.pk}/?rp={page_no}"))
        top = text.split("Все пометки")[0]
        assert "загадка-999" in top  # нераспознанная - всегда сверху отдельным списком
        assert "Начало без конца" in top  # непарная - тоже
        assert f"Страница {page_no} из 3" in text


def test_issues_with_cells_shown(client, users, login, tmp_path, post_new):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append(["Дистанция, м", "U т-з, В", "U т-з выкл, В"])
    ws.append([0, -1.2, -0.9])
    ws.append([10, "текст", -0.9])
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    login(users["a"])
    post_new(path)
    worker.process_all()
    job = MapJob.objects.get()
    assert job.state == MapJob.State.CHECK_FAILED
    page = _content(client.get(f"/maps/{job.pk}/"))
    assert "<code>B3</code>" in page
    assert json.dumps(job.check_result)  # результат проверки - обычный JSON


def test_split_form_km_messages_hint_and_current_bounds(client, users, login, small_book, post_new):
    login(users["a"])
    post_new(small_book, start_km="12,5")
    job = MapJob.objects.get()
    assert worker.process_all() == 1
    page = _content(client.get(f"/maps/{job.pk}/"))
    # пример по фактическому километражу (6 км от 12,5), а не абстрактное «0; 3; 5,5»
    assert "12,5; 15,5; 18,5" in page and "0; 3; 5,5" not in page
    # режим «авто»: поле пустое, текущие границы - отдельной строкой
    assert "Текущие границы:" in page and "<code>12,5; 18,5</code>" in page
    url = f"/maps/{job.pk}/split/"
    bad = {
        ("bounds", "15; 13; 18"): "15 -&gt; 13 км",
        ("bounds", "15"): "12,5; 15,5; 18,5",
        ("length", "0"): "положительная длина",
        ("length", "0,01"): "(не меньше 0,02 км)",
    }
    for (mode, value), text in bad.items():
        response = client.post(url, {"mode": mode, "value": value})
        assert response.status_code == 400, (mode, value)
        content = _content(response)
        assert text in content, (mode, value)
        assert "2500 -" not in content
