"""Удаление, журнал, админка, настройки, демо-пользователи."""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import pytest
from config.dburl import parse_database_url
from django.core.management import CommandError, call_command
from maps import storage
from maps.models import AuditLog, MapJob

ROOT = Path(__file__).resolve().parents[2]


# --- удаление ---------------------------------------------------------------------------


def test_delete_removes_files_and_record(client, users, login, done_job):
    job = done_job(users["a"])
    source = storage.resolve(job.source_path)
    results = storage.resolve(storage.job_results_rel(job.pk))
    assert source.is_file() and results.is_dir()
    login(users["a"])
    page = client.get(f"/maps/{job.pk}/delete/")
    assert page.status_code == 200
    assert "Да, удалить" in page.content.decode()
    assert MapJob.objects.filter(pk=job.pk).exists()  # GET ничего не удаляет
    response = client.post(f"/maps/{job.pk}/delete/")
    assert response.status_code == 302 and response["Location"] == "/"
    assert not MapJob.objects.filter(pk=job.pk).exists()
    assert not source.exists()
    assert not results.exists()
    entry = AuditLog.objects.get(action="delete")
    assert entry.username == "alice" and "Тестовая карта" in entry.object_repr
    assert entry.job is None  # карта удалена, запись журнала осталась


def test_cannot_delete_foreign_map(client, users, login, done_job):
    job = done_job(users["a"])
    login(users["b"])
    assert client.post(f"/maps/{job.pk}/delete/").status_code == 404
    assert MapJob.objects.filter(pk=job.pk).exists()
    assert storage.resolve(job.pdf_path).is_file()
    assert not AuditLog.objects.filter(action="delete").exists()


def test_can_delete_while_building(client, users, login, make_job):
    """Карта удаляется и в работе (ТЗ, раздел 9); подробнее - test_web_cancel.py."""
    job = make_job(users["a"], state=MapJob.State.BUILDING)
    login(users["a"])
    response = client.post(f"/maps/{job.pk}/delete/")
    assert response.status_code == 302
    assert not MapJob.objects.filter(pk=job.pk).exists()


def test_delete_page_has_no_js_confirm(client, users, login, make_job):
    job = make_job(users["a"], check=True)
    login(users["a"])
    for url in ("/", f"/maps/{job.pk}/", f"/maps/{job.pk}/delete/"):
        html = client.get(url).content.decode()
        assert "confirm(" not in html and "alert(" not in html and "onclick" not in html


def test_pages_have_no_external_resources(client, users, login, make_job):
    job = make_job(users["a"], check=True)
    login(users["a"])
    for url in ("/", "/maps/new/", f"/maps/{job.pk}/"):
        html = client.get(url).content.decode()
        assert "http://" not in html and "https://" not in html and "//cdn" not in html


def test_admin_deletion_in_admin_removes_files(users, make_job):
    job = make_job(users["a"])
    source = storage.resolve(job.source_path)
    MapJob.objects.get(pk=job.pk).delete()
    assert not source.exists()


# --- журнал ------------------------------------------------------------------------------


def test_audit_records_ip(client, users, login, small_book, post_new):
    login(users["a"])
    client.post("/maps/new/", {}, REMOTE_ADDR="192.0.2.7")
    post_new(small_book)
    entry = AuditLog.objects.filter(action="upload").get()
    assert entry.user == users["a"] and entry.ip == "127.0.0.1"
    assert entry.job is not None and "small.xlsx" in entry.details


def test_audit_download_logged(client, users, login, done_job):
    job = done_job(users["a"])
    login(users["admin"])
    client.get(f"/maps/{job.pk}/files/pdf/", REMOTE_ADDR="198.51.100.4")
    entry = AuditLog.objects.get(action="download")
    assert entry.username == "root" and entry.ip == "198.51.100.4"
    assert entry.job == job


# --- админка -------------------------------------------------------------------------------


def test_admin_pages(client, users, login, done_job):
    from django.contrib.auth import get_user_model

    User = get_user_model()
    User.objects.create_superuser("super", "s@example.com", "Karta-2026-test")
    client.force_login(User.objects.get(username="super"))
    job = done_job(users["a"])
    AuditLog.objects.create(action="login", username="alice")
    assert client.get("/admin/").status_code == 200
    changelist = client.get("/admin/maps/mapjob/?state__exact=done")
    assert changelist.status_code == 200 and job.object_name in changelist.content.decode()
    assert client.get(f"/admin/maps/mapjob/?owner__id__exact={users['a'].pk}").status_code == 200
    assert client.get("/admin/maps/auditlog/").status_code == 200
    # журнал только для чтения
    assert client.get("/admin/maps/auditlog/add/").status_code == 403
    entry = AuditLog.objects.first()
    assert (
        client.post(f"/admin/maps/auditlog/{entry.pk}/delete/", {"post": "yes"}).status_code == 403
    )
    assert client.post(f"/admin/maps/auditlog/{entry.pk}/change/", {}).status_code == 403
    assert client.get("/admin/auth/user/").status_code == 200
    assert client.get(f"/admin/maps/mapjob/{job.pk}/change/").status_code == 200


def test_admin_forbidden_for_regular_user(client, users, login):
    login(users["a"])
    response = client.get("/admin/maps/mapjob/")
    assert response.status_code in (302, 403)
    if response.status_code == 302:
        assert "/login/" in response["Location"]


# --- настройки -------------------------------------------------------------------------------


def _import_settings(**env):
    full = {k: v for k, v in os.environ.items() if not k.startswith("CORMAP_")}
    full.update(env)
    full["PYTHONPATH"] = str(ROOT / "web")
    full["PYTHONUTF8"] = "1"  # иначе текст ошибки придёт в кодировке консоли Windows
    return subprocess.run(
        [sys.executable, "-c", "import config.settings as s; print(s.DEBUG, s.SECRET_KEY[:4])"],
        env=full,
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )


def test_secret_key_required_without_debug():
    result = _import_settings(CORMAP_DEBUG="0")
    assert result.returncode != 0
    assert "CORMAP_SECRET_KEY" in result.stderr


def test_production_settings_with_secret(tmp_path):
    result = _import_settings(
        CORMAP_DEBUG="0", CORMAP_SECRET_KEY="abcdefgh", CORMAP_DATA_DIR=str(tmp_path / "d")
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.split() == ["False", "abcd"]


def test_local_mode_is_explicit(tmp_path):
    result = _import_settings(CORMAP_DEBUG="1", CORMAP_DATA_DIR=str(tmp_path / "d"))
    assert result.returncode == 0, result.stderr
    assert result.stdout.startswith("True")


@pytest.mark.parametrize("value", [None, "", "0", "debug", "disabled", "2"])
def test_debug_is_off_unless_explicitly_enabled(tmp_path, value):
    """Без явного CORMAP_DEBUG=1 отладка выключена: без ключа сайт не запускается."""
    env = {"CORMAP_DATA_DIR": str(tmp_path / "d")}
    if value is not None:
        env["CORMAP_DEBUG"] = value
    result = _import_settings(**env)
    assert result.returncode != 0
    assert "CORMAP_SECRET_KEY" in result.stderr
    with_key = _import_settings(**env, CORMAP_SECRET_KEY="abcdefgh")
    assert with_key.returncode == 0, with_key.stderr
    assert with_key.stdout.split() == ["False", "abcd"]


def test_parse_database_url():
    pg = parse_database_url("postgres://cormap:p%40ss@db.local:5433/cormapdb")
    assert pg == {
        "ENGINE": "django.db.backends.postgresql",
        "NAME": "cormapdb",
        "USER": "cormap",
        "PASSWORD": "p@ss",
        "HOST": "db.local",
        "PORT": "5433",
        "CONN_MAX_AGE": 60,
        "CONN_HEALTH_CHECKS": True,
    }
    assert parse_database_url("postgresql://u@h/db")["PORT"] == ""
    assert parse_database_url("sqlite:///var/db.sqlite3")["NAME"] == "var/db.sqlite3"
    assert parse_database_url("sqlite:////abs/db.sqlite3")["NAME"] == "/abs/db.sqlite3"
    for bad in ("mysql://u@h/db", "postgres://u@h/", "sqlite://"):
        with pytest.raises(ValueError):
            parse_database_url(bad)


def test_settings_basics(settings):
    assert settings.LANGUAGE_CODE == "ru"
    assert settings.TIME_ZONE == "Europe/Moscow"
    assert len(settings.AUTH_PASSWORD_VALIDATORS) >= 4
    assert settings.CORMAP_MAX_UPLOAD_BYTES == 50 * 1024 * 1024
    assert settings.DATA_UPLOAD_MAX_MEMORY_SIZE < settings.CORMAP_MAX_UPLOAD_BYTES
    assert "django.contrib.auth.middleware.LoginRequiredMiddleware" in settings.MIDDLEWARE


# --- демо-пользователи ----------------------------------------------------------------------


def test_seed_demo(db, settings, capsys):
    from django.contrib.auth import get_user_model

    settings.DEBUG = True
    call_command("seed_demo")
    User = get_user_model()
    assert User.objects.filter(username__in=["ivanov", "petrov", "admin"]).count() == 3
    assert User.objects.get(username="admin").is_superuser
    assert not User.objects.get(username="ivanov").is_staff
    assert User.objects.get(username="petrov").check_password("Karta-2026-test")
    capsys.readouterr()
    call_command("seed_demo")  # повтор ничего не создаёт
    assert "Уже есть: ivanov" in capsys.readouterr().out
    assert User.objects.count() == 3


def test_seed_demo_refuses_without_debug(db, settings):
    settings.DEBUG = False
    with pytest.raises(CommandError):
        call_command("seed_demo")


def test_worker_module_importable_without_django_in_core():
    code = (
        "import sys, cormap.build, cormap.check, cormap.serialize;"
        "assert 'django' not in sys.modules, 'ядро импортирует Django'"
    )
    result = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


# --- шаблон книги (ТЗ, 4.6) ------------------------------------------------------------------


def test_template_requires_login(client):
    response = client.get("/maps/template/")
    assert response.status_code == 302 and "/login" in response["Location"]


def test_template_download(client, users, login):
    from io import BytesIO

    from openpyxl import load_workbook

    login(users["a"])
    response = client.get("/maps/template/")
    assert response.status_code == 200
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "shablon_knigi.xlsx" in response["Content-Disposition"]
    assert "attachment" in response["Content-Disposition"]
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames[0] == "Template" and "Объекты" in wb.sheetnames


def test_template_sheet_with_refinery(client, users, login, fake_refinery):
    from io import BytesIO

    from openpyxl import load_workbook

    login(users["a"])
    wb = load_workbook(BytesIO(client.get("/maps/template/").content))
    assert wb.sheetnames[0] == "Template"


def test_refinery_script_only_with_module(client, users, login):
    """Без модуля блока обработки и его скрипта на форме нет."""
    login(users["a"])
    assert "maps/refinery.js" not in client.get("/maps/new/").content.decode()


def test_refinery_script_on_forms(client, users, login, fake_refinery):
    """Блокировка несочетаемых флажков (ТЗ, 4.5) подключена отдельным файлом (CSP)."""
    login(users["a"])
    page = client.get("/maps/new/").content.decode()
    assert 'id="refinery-params"' in page and "maps/refinery.js" in page


def test_template_link_on_new_page(client, users, login):
    login(users["a"])
    page = client.get("/maps/new/").content.decode()
    assert "Скачать шаблон книги" in page and "/maps/template/" in page
