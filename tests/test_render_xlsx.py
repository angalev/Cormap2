"""Тесты книги Excel с графиками (этап 7)."""

from __future__ import annotations

import re
import time

import pytest
from openpyxl import load_workbook

from cormap.build import build_map
from cormap.check import check_workbook
from cormap.cli import main
from cormap.defects import defect_objects
from cormap.layout import split_manual
from cormap.params import BuildParams
from cormap.render_pdf import (
    STRIP_INSULATION,
    STRIP_OBJECTS,
    STRIP_SOIL,
    make_context,
    render_strip_png,
    strip_image_size_mm,
)
from cormap.render_xlsx import (
    SHEET_DEFECTS,
    SHEET_KM,
    SHEET_MAP,
    SHEET_NORMS,
    write_xlsx,
)
from tests.conftest import BASE
from tests.synth import HEADERS_V2, make_sample_like, make_workbook

NAME = "Тестовый объект"
REF = re.compile(
    r"^'?(?P<sheet>[^'!]+)'?!\$(?P<c1>[A-Z]+)\$(?P<r1>\d+):\$(?P<c2>[A-Z]+)\$(?P<r2>\d+)$"
)

# графики листа 1 и листа 2 (по порядку на участке) и заголовки
T_POTENTIAL = "Потенциал трубопровода в грунте, В"
T_GRADIENT = "Градиент потенциала, мВ"
T_DEFECTS = "Площадь дефекта"
T_ZONES = "Зоны коррозионной опасности"
T_DEPTH = "Глубина залегания трубопровода, м"
T_RES = "Удельное электрическое сопротивление грунта, Ом·м"


@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    return make_sample_like(tmp_path_factory.mktemp("x") / "sample.xlsx")


def make_xlsx(source, out, split="0;3000;5500", params=None):
    """Строит только книгу Excel (без PDF): те же шаги, что в build_map."""
    survey, report, notes = check_workbook(source)
    assert report.ok, report.format_text()
    params = params or BuildParams(NAME, split=split)
    objects = (*notes.objects, *defect_objects(survey))
    lo, hi = survey.distance_range
    sections = split_manual(lo, hi, params.split, None, objects, params.paper).sections
    ctx = make_context(survey, notes.objects, len(sections))
    write_xlsx(out, source, survey, sections, objects, params, ctx)
    return sections, survey, ctx, params


def chart_title(chart) -> str:
    return " ".join(r.t for p in chart.title.tx.rich.p for r in p.r)


def parse_ref(text: str):
    m = REF.match(text)
    assert m, text
    return m["sheet"], m["c1"], int(m["r1"]), m["c2"], int(m["r2"])


def source_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        return [_trim(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def sheet_rows(path, name):
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[name]
        ws.reset_dimensions()
        return [_trim(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _trim(row):
    row = list(row)
    while row and row[-1] is None:
        row.pop()
    return row


# --- лист «Данные» -------------------------------------------------------------


def test_sheets_present_and_data_equal_sample(sample, tmp_path):
    out = tmp_path / "k.xlsx"
    make_xlsx(sample, out)
    wb = load_workbook(out, read_only=True)
    assert {"Result", SHEET_MAP, SHEET_DEFECTS, SHEET_NORMS, SHEET_KM} <= set(wb.sheetnames)
    wb.close()
    src, dst = source_rows(sample), sheet_rows(out, "Result")
    assert src[0] == HEADERS_V2 and dst[0] == HEADERS_V2
    assert len(dst) == len(src) and dst == src  # заголовки и все значения по ячейкам
    # заголовок закреплён
    wb = load_workbook(out)
    assert wb["Result"].freeze_panes == "A2"


def test_data_equal_permuted_columns_and_extra_column(book, tmp_path):
    cols = [
        "Комментарий",
        "Лишняя колонка",
        "U т-з выкл, В",
        "Дистанция, м",
        "U т-з, В",
        "Дистанция грун, м",
        "GPS",
        "Уд. соп-е грунтов, Ом*м",
        "Тип грунта",
        "Дистанция глубины, м",
        "Глубина заложения трубы, м",
    ]
    rows = []
    for i in range(0, 200, 2):
        rows.append(
            [
                "КИП, КИП-1" if i == 10 else None,
                f"заметка {i}" if i % 20 == 0 else i,
                -1.3,
                i,
                -1.5 - i / 1000,
                float(i) if i % 10 == 0 else None,
                "55.1;37.2" if i == 0 else None,
                100 + i if i % 10 == 0 else None,
                "песчаный" if i == 0 else None,
                float(i) if i % 20 == 0 else None,
                1.2 if i % 20 == 0 else None,
            ]
        )
    rows.insert(7, [None] * len(cols))  # пустая строка посередине
    path = book(cols, rows)
    out = tmp_path / "k.xlsx"
    make_xlsx(path, out, split="100")
    src, dst = source_rows(path), sheet_rows(out, "Result")
    assert dst == src and dst[0] == cols
    # ничего не потеряно: нераспознанная колонка и GPS на месте
    assert any(r[1] == "заметка 0" for r in dst[1:] if len(r) > 1)


def test_data_equal_values_roundtrip_types(book, tmp_path):
    rows = [[0, -1.5, -1.3], [2, "-1,4", -1.2], [4, -1.4, None], [6, -1.4, -1.2]]
    path = book(BASE, rows)
    out = tmp_path / "k.xlsx"
    make_xlsx(path, out, split="6")
    assert sheet_rows(out, "Result") == source_rows(path)


# --- лист «Карта»: диаграммы ------------------------------------------------------


@pytest.fixture(scope="module")
def built(sample, tmp_path_factory):
    out = tmp_path_factory.mktemp("k") / "k.xlsx"
    sections, survey, ctx, params = make_xlsx(sample, out)
    wb = load_workbook(out)
    return out, wb, sections, survey


def test_chart_and_image_counts(built):
    _, wb, sections, _ = built
    ws = wb[SHEET_MAP]
    assert len(sections) == 2
    # на участке 7 графиков (потенциал, 2 градиента, дефекты, зоны, глубина, УЭС)
    assert len(ws._charts) == 7 * len(sections)
    assert len(ws._images) == 3 * len(sections)  # изоляция, грунт, объекты
    titles = [chart_title(c) for c in ws._charts[:7]]
    assert titles == [T_POTENTIAL, T_GRADIENT, T_GRADIENT, T_DEFECTS, T_ZONES, T_DEPTH, T_RES]


def test_series_refer_to_section_rows(built):
    out, wb, sections, _ = built
    data = sheet_rows(out, "Result")
    km_to_data = {"A": 0, "B": HEADERS_V2.index("Дистанция грун, м"),
                  "C": HEADERS_V2.index("Дистанция глубины, м")}  # fmt: skip
    ws = wb[SHEET_MAP]
    checked = 0
    for k, sec in enumerate(sections):
        for chart in ws._charts[7 * k : 7 * k + 7]:
            for s in chart.series:
                ysheet, yc1, y1, yc2, y2 = parse_ref(s.yVal.numRef.f)
                xsheet, xc1, x1, xc2, x2 = parse_ref(s.xVal.numRef.f)
                assert xc1 == xc2 and yc1 == yc2  # один столбец
                if ysheet != "Result":
                    # линия норматива: две строки служебного листа
                    assert ysheet == SHEET_NORMS and xsheet == SHEET_NORMS
                    continue
                assert xsheet == SHEET_KM  # ось X - километраж, не дублируется на «Данных»
                assert (x1, x2) == (y1, y2)  # те же строки
                dcol = km_to_data[xc1]
                inside = [
                    float(data[r - 1][dcol])
                    for r in range(y1, y2 + 1)
                    if len(data[r - 1]) > dcol and data[r - 1][dcol] is not None
                ]
                a, b = sec.start_m, sec.end_m
                # внутри диапазона - точки участка и по одной соседней с каждой стороны
                assert min(inside) <= a and max(inside) >= b or (y1 == 2 or y2 == len(data))
                core = [v for v in inside if a <= v <= b]
                assert len(inside) - len(core) <= 2
                assert core and all(a <= v <= b for v in core)
                # соседей ровно столько, сколько не помещается в участок
                assert inside[0] < a or inside[0] == core[0]
                assert inside[-1] > b or inside[-1] == core[-1]
                checked += 1
    assert checked >= 2 * 8


def test_axes_and_style(built):
    _, wb, sections, _ = built
    ws = wb[SHEET_MAP]
    for k, sec in enumerate(sections):
        charts = ws._charts[7 * k : 7 * k + 7]
        pot = charts[0]
        # ось потенциала перевёрнута: сверху -3,00, снизу -0,50, шаг 0,5
        assert pot.y_axis.scaling.orientation == "maxMin"
        assert (pot.y_axis.scaling.min, pot.y_axis.scaling.max) == (-3.0, -0.5)
        assert pot.y_axis.majorUnit == 0.5
        assert len(pot.series) == 4
        for c in charts:
            assert c.legend.position == "b"
            assert c.x_axis.number_format.formatCode == "0.00"
            assert c.x_axis.scaling.min == pytest.approx(sec.start_m / 1000)
            assert c.x_axis.scaling.max == pytest.approx(
                (sec.start_m + sec.nominal_length_m) / 1000
            )
            assert c.x_axis.majorUnit == pytest.approx(sec.tick_step_m / 1000)
            assert c.x_axis.title is not None and c.y_axis.title is not None
            for s in c.series:
                assert s.marker.symbol is None and not s.smooth  # "none" читается как None
        # пунктир нормативов и цвета линий потенциала
        colors = [s.graphicalProperties.line.solidFill.srgbClr for s in pot.series]
        assert colors == ["00B050", "00B0F0", "1F3F7F", "FF0000"]
        dashes = [s.graphicalProperties.line.prstDash for s in pot.series]
        assert dashes == ["solid", "solid", "dash", "dash"]
        # площадь дефекта 0-3, зоны 0-1, глубина с нормативом
        assert (charts[3].y_axis.scaling.min, charts[3].y_axis.scaling.max) == (0, 3)
        assert (charts[4].y_axis.scaling.min, charts[4].y_axis.scaling.max) == (0, 1)
        assert len(charts[5].series) == 2 and len(charts[6].series) == 3


def test_chart_and_image_share_width(built):
    _, wb, _, _ = built
    ws = wb[SHEET_MAP]
    widths = {c.anchor.ext.width for c in ws._charts} | {i.anchor.ext.width for i in ws._images}
    assert len(widths) == 1
    # картинки и диаграммы идут сверху вниз без наложений внутри страницы
    for chart_or_img in (*ws._charts, *ws._images):
        assert chart_or_img.anchor._from.col == 0


def test_service_sheets(built):
    out, wb, sections, _ = built
    norms = wb[SHEET_NORMS]
    assert norms["A1"].value and norms["B1"].value == 0.0
    # две точки на участок: начало и конец номинальной длины
    assert [norms.cell(row=r, column=2).value for r in (4, 5, 6, 7)] == [
        0.0,
        3000.0,
        3000.0,
        5500.0,
    ]
    assert norms["D4"].value == -2.5 and norms["E4"].value == -0.85 and norms["F4"].value == 0.8
    assert (norms["G4"].value, norms["H4"].value) == (20.0, 50.0)
    # километраж - формула от листа «Данные» (лист «Данные» не меняется)
    km = wb[SHEET_KM]
    assert km["A2"].value == "='Result'!A2/1000+'Нормативы'!$B$1"
    assert km["B2"].value.startswith("='Result'!J2/1000")
    # «Площадь дефекта»: ступеньки точечных повреждений видны на служебном листе
    classes = {norms.cell(row=r, column=14).value for r in range(4, 40)}
    assert {0, 1, 2, 3} <= classes


def test_start_km_in_axes_and_norms(sample, tmp_path):
    out = tmp_path / "k.xlsx"
    params = BuildParams(NAME, start_km=12.5, split="0;3000;5500")
    make_xlsx(sample, out, params=params)
    wb = load_workbook(out)
    assert wb[SHEET_NORMS]["B1"].value == 12.5
    ws = wb[SHEET_MAP]
    c = ws._charts[0]
    assert c.x_axis.scaling.min == 12.5 and c.x_axis.scaling.max == 15.5
    assert "12,50 - 15,50" in ws["A2"].value or "12,50 - 15,00" in ws["A2"].value


def test_res_axis_max_in_excel(sample, tmp_path):
    out = tmp_path / "k.xlsx"
    params = BuildParams(NAME, split="0;3000;5500", res_axis_max=130)
    make_xlsx(sample, out, params=params)
    ws = load_workbook(out)[SHEET_MAP]
    res_charts = [c for c in ws._charts if chart_title(c) == T_RES]
    assert len(res_charts) == 2
    for c in res_charts:
        assert (c.y_axis.scaling.min, c.y_axis.scaling.max) == (0.0, 130.0)


# --- старый шаблон ------------------------------------------------------------------


def test_old_template_has_no_data_lines(book, tmp_path):
    rows = [[i * 2, -1.5, -1.3] for i in range(100)]
    path = book(BASE, rows)
    out = tmp_path / "k.xlsx"
    make_xlsx(path, out, split="198")
    wb = load_workbook(out)
    ws = wb[SHEET_MAP]
    # остаются потенциал, площадь дефекта и зоны
    assert [chart_title(c) for c in ws._charts] == [T_POTENTIAL, T_DEFECTS, T_ZONES]
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    nodata = [t for t in texts if t.endswith(": нет данных")]
    assert sorted(nodata) == sorted(
        [
            f"{T_GRADIENT} (левый): нет данных",
            f"{T_GRADIENT} (правый): нет данных",
            f"{T_DEPTH}: нет данных",
            f"{T_RES}: нет данных",
        ]
    )
    assert len(ws._images) == 3  # ленты рисуются и без данных («нет данных» на картинке)
    assert wb[SHEET_DEFECTS]["A2"].value == "Повреждений нет"


# --- печать -------------------------------------------------------------------------


def test_print_setup(built):
    _, wb, sections, _ = built
    ws = wb[SHEET_MAP]
    ps = ws.page_setup
    assert ps.orientation == "landscape" and ps.paperSize == 9  # A4
    assert ps.fitToWidth == 1 and ps.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    pages = 2 * len(sections)
    breaks = [b.id for b in ws.row_breaks.brk]
    assert len(breaks) == pages - 1
    step = breaks[0]
    assert breaks == [step * k for k in range(1, pages)]
    assert ws.print_area.endswith(f"${pages * step}")
    assert ws.oddFooter.center.text == "Лист &P из &N"
    # листы 1 и 2 участка начинаются с заголовка на своей странице
    for k in range(pages):
        assert ws.cell(row=k * step + 1, column=1).value == NAME
    # диаграммы страницы не выходят за её строки
    for chart in ws._charts:
        page = chart.anchor._from.row // step
        end_row = chart.anchor._from.row + (chart.anchor.ext.height / 190500)
        assert end_row <= (page + 1) * step + 1e-6


def test_print_a3(sample, tmp_path):
    out = tmp_path / "k.xlsx"
    make_xlsx(sample, out, params=BuildParams(NAME, split="5500", paper="a3"))
    ws = load_workbook(out)[SHEET_MAP]
    assert ws.page_setup.paperSize == 8 and ws.page_setup.orientation == "landscape"


# --- лист «Повреждения» --------------------------------------------------------------


def test_defects_sheet(sample, tmp_path):
    out = tmp_path / "k.xlsx"
    make_xlsx(sample, out)
    rows = sheet_rows(out, SHEET_DEFECTS)
    assert rows[0][0] == "Начало, м" and rows[0][-1] == "Строка книги"
    body = rows[1:]
    assert len(body) == 5  # 3 из блока «Повреждения» и 2 пометки «ПИП»
    assert all(r[9] == "ручной" for r in body)
    by_start = {r[0]: r for r in body}
    assert [by_start[d][8] for d in (800.0, 2600.0, 4000.0)] == [1, 2, 3]  # итоговый класс
    # строка книги: блок «Повреждения» идёт с 2-й строки; пометки «ПИП» - строки комментариев
    assert [by_start[d][10] for d in (800.0, 2600.0, 4000.0)] == [2, 3, 4]
    pip = [r for r in body if r[0] in (1800.0, 3500.0)]
    assert len(pip) == 2 and all(r[10] and r[10] > 100 for r in pip)
    assert all(r[2] == r[0] / 1000 for r in body)  # километраж при start_km = 0


# --- build_map и командная строка -------------------------------------------------


def test_build_map_writes_xlsx(sample, tmp_path):
    res = build_map(sample, BuildParams(NAME, split="0;3000;5500"), tmp_path, png=False)
    assert res.xlsx_path == tmp_path / "karta.xlsx" and res.xlsx_path.exists()
    assert "Excel: графики и картинки" in res.timings and "Excel: запись файла" in res.timings
    wb = load_workbook(res.xlsx_path)
    assert len(wb[SHEET_MAP]._charts) == 14
    res = build_map(sample, BuildParams(NAME, split="5500"), tmp_path / "n", png=False, xlsx=False)
    assert res.xlsx_path is None and not (tmp_path / "n" / "karta.xlsx").exists()


def test_cli_no_xlsx(sample, tmp_path, capsys):
    out = tmp_path / "o"
    assert main(["build", str(sample), "--out", str(out), "--name", NAME, "--split", "5500",
                 "--no-png"]) == 0  # fmt: skip
    assert (out / "karta.xlsx").exists() and "Excel:" in capsys.readouterr().out
    out2 = tmp_path / "o2"
    assert main(["build", str(sample), "--out", str(out2), "--name", NAME, "--split", "5500",
                 "--no-png", "--no-xlsx"]) == 0  # fmt: skip
    assert (out2 / "karta.pdf").exists() and not (out2 / "karta.xlsx").exists()


# --- картинки лент ----------------------------------------------------------------------


def test_strip_png_size_matches_geometry(sample):
    from io import BytesIO

    from PIL import Image

    from cormap.layout import place_objects
    from cormap.render_pdf import SheetContext  # noqa: F401

    survey, report, notes = check_workbook(sample)
    params = BuildParams(NAME, split="5500")
    objects = (*notes.objects, *defect_objects(survey))
    sec = split_manual(0, 5500, 5500.0, None, objects).sections[0]
    ctx = make_context(survey, notes.objects, 1)
    placement = place_objects(objects, sec)
    for kind in (STRIP_INSULATION, STRIP_SOIL, STRIP_OBJECTS):
        png = render_strip_png(kind, sec, placement, params, ctx, dpi=100)
        w_mm, h_mm = strip_image_size_mm(kind, params, ctx, placement)
        with Image.open(BytesIO(png)) as im:
            assert im.format == "PNG"
            assert im.width == pytest.approx(w_mm / 25.4 * 100, abs=1.5)
            assert im.height == pytest.approx(h_mm / 25.4 * 100, abs=1.5)
    # ширина картинки = ширина листа без полей (плюс поле картинки)
    assert strip_image_size_mm(STRIP_SOIL, params, ctx, placement)[0] == pytest.approx(278.0)
    with pytest.raises(ValueError):
        strip_image_size_mm("abc", params, ctx, placement)


# --- тяжёлые ------------------------------------------------------------------------------


@pytest.mark.slow
def test_build_50km_with_excel_under_60s(tmp_path, capsys):
    path = make_workbook(tmp_path / "b.xlsx", length_m=50_000, step_m=2.0)
    t0 = time.perf_counter()
    res = build_map(path, BuildParams(NAME), tmp_path / "out")
    elapsed = time.perf_counter() - t0
    size = res.xlsx_path.stat().st_size / 1e6
    with capsys.disabled():
        print(
            f"\n[slow] 50 км (PDF+PNG+Excel): {elapsed:.1f} с, участков {len(res.sections)}, "
            f"karta.xlsx {size:.1f} МБ; {res.format_timings()}"
        )
    assert elapsed <= 60
    wb = load_workbook(res.xlsx_path, read_only=True)
    assert SHEET_MAP in wb.sheetnames
    wb.close()


@pytest.mark.slow
def test_build_300km_with_excel_under_5min(tmp_path, capsys):
    notes = [(d, "КИП, КИП") for d in range(0, 300_000, 25_000)]
    path = make_workbook(tmp_path / "b.xlsx", length_m=300_000, step_m=2.0, notes=notes)
    t0 = time.perf_counter()
    res = build_map(path, BuildParams(NAME), tmp_path / "out")
    elapsed = time.perf_counter() - t0
    size = res.xlsx_path.stat().st_size / 1e6
    with capsys.disabled():
        print(
            f"\n[slow] 300 км (PDF+PNG+Excel): {elapsed:.1f} с, участков {len(res.sections)}, "
            f"karta.xlsx {size:.1f} МБ; {res.format_timings()}"
        )
    assert len(res.sections) == 3 and res.xlsx_path.exists()
    assert elapsed <= 300
