"""Заглушка модуля `refinery` для тестов (настоящий модуль пишет заказчик).

Повторяет интерфейс из ТЗ, раздел 4.5: читает лист `Template`, пишет лист `Result` с теми же
заголовками и строками. Режимы: ok, fail (success=False), raise (исключение), no_file
(«успех» без файла). Заглушка подкладывается в sys.modules фикстурой `fake_refinery`.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from cormap.preprocess import FLAG_NAMES


@dataclass
class RefinerySettings:
    utz: bool = False
    utzd: bool = False
    utzdc: bool = False
    grlu: bool = False
    grld: bool = False
    grluc: bool = False
    grldc: bool = False
    grru: bool = False
    grrd: bool = False
    grruc: bool = False
    grrdc: bool = False
    step: int = 5
    random_seed: int | None = None


@dataclass
class RefineryStats:
    input_rows: int
    output_rows: int
    distance_range_m: tuple[float, float]
    columns_generated: list[str]
    calculation_time_ms: int


@dataclass
class RefineryResult:
    success: bool
    output_file: Path | None = None
    error_code: str | None = None
    error_message: str | None = None
    stats: RefineryStats | None = None


@dataclass
class FakeState:
    mode: str = "ok"
    calls: list[dict[str, Any]] = field(default_factory=list)


STATE = FakeState()


def run_refinery(
    input_file: Path | bytes, settings: RefinerySettings, output_file: Path | None = None
) -> RefineryResult:
    STATE.calls.append({"input_file": input_file, "settings": settings, "output_file": output_file})
    if STATE.mode == "raise":
        raise RuntimeError("внезапный сбой заглушки")
    if STATE.mode == "fail":
        return RefineryResult(False, None, "bad_input", "Нет данных для обработки.", None)
    wb = load_workbook(input_file, read_only=True, data_only=True)
    if "Template" not in wb.sheetnames:
        return RefineryResult(False, None, "no_template", "Нет листа Template.", None)
    rows = [list(r) for r in wb["Template"].iter_rows(values_only=True)]
    wb.close()
    if STATE.mode == "no_file":
        return RefineryResult(True, output_file, None, None, RefineryStats(0, 0, (0, 0), [], 1))
    out = Workbook()
    ws = out.active
    ws.title = "Result"
    for row in rows:
        ws.append(row)
    out.save(output_file)
    dists = [r[0] for r in rows[1:] if isinstance(r[0], int | float)]
    generated = [name for name in FLAG_NAMES if getattr(settings, name)]
    stats = RefineryStats(
        input_rows=len(rows) - 1,
        output_rows=len(rows) - 1,
        distance_range_m=(min(dists), max(dists)) if dists else (0.0, 0.0),
        columns_generated=generated,
        calculation_time_ms=12,
    )
    return RefineryResult(True, output_file, None, None, stats)


@pytest.fixture
def fake_refinery(monkeypatch):
    """Подкладывает заглушку как `refinery`; возвращает состояние (mode, calls)."""
    STATE.mode = "ok"
    STATE.calls = []
    monkeypatch.setitem(sys.modules, "refinery", sys.modules[__name__])
    return STATE
