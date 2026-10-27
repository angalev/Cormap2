"""Шаблон книги (ТЗ, 4.6): состав, справочник шифров, проверка данных, чтение заполненной книги."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from cormap.cli import main
from cormap.model import INSULATION_TYPES, SOIL_TYPES, CommentItem
from cormap.notes import OBJECT_TYPES, POINT, parse_comments
from cormap.reader import COLUMNS, read_workbook
from cormap.template import TEMPLATE_KEYS, build_template, list_entries
from tests.conftest import REFERENCE


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "tpl.xlsx"
    build_template(path)
    return path


def _validations(ws):
    return {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}


def test_keys_match_columns():
    assert set(TEMPLATE_KEYS) == set(COLUMNS)
    assert len(TEMPLATE_KEYS) == len(COLUMNS)


def test_headers_and_order(book):
    ws = load_workbook(book)["Result"]
    headers = [c.value for c in ws[1]]
    assert headers == [COLUMNS[k][0] for k in TEMPLATE_KEYS]
    ref = load_workbook(REFERENCE / "template.xlsx").worksheets[0]
    ref_headers = [c.value for c in ref[1] if c.value is not None]
    assert headers[:11] == ref_headers[:11]
    assert headers[11] == "Тип грунта"  # сразу за УЭС: модуль обработки ищет колонки по номеру
    assert headers[12:16] == ref_headers[11:15]
    assert headers[16:] == ["Дистанция изоляции, м", "Тип изоляции"]
    assert ws.max_row == 1
    assert ws.freeze_panes == "A2"


def test_sheet_names_and_order(tmp_path):
    path = tmp_path / "t.xlsx"
    build_template(path, "Template")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Template", "Объекты", "Как заполнять", "Списки"]
    assert wb["Списки"].sheet_state == "hidden"
    assert wb.active.title == "Template"
    assert 10 <= wb["Как заполнять"].max_row <= 15


def test_every_entry_recognized_as_source(tmp_path):
    entries = list_entries()
    items = [CommentItem(float(i), i + 2, e.text) for i, e in enumerate(entries)]
    for item, entry in zip(items, entries, strict=True):
        result = parse_comments([item])
        assert len(result.records) == 1, entry.text
        rec = result.records[0]
        assert (rec.type_key, rec.role) == (entry.type_key, entry.role), entry.text


def test_every_type_present():
    present = {e.type_key for e in list_entries()}
    assert present == {t.key for t in OBJECT_TYPES}
    for t in OBJECT_TYPES:
        if t.extended:
            roles = {e.role for e in list_entries() if e.type_key == t.key}
            assert len(roles) == 3
    assert "болото" in {e.text for e in list_entries()}
    assert "болот" not in {e.text for e in list_entries()}
    assert all(e.role == POINT for e in list_entries() if e.type_key == "pit")


def test_lists_sheet_and_objects_sheet(book):
    wb = load_workbook(book)
    lists = wb["Списки"]
    texts = [e.text for e in list_entries()]
    assert [lists.cell(row=r, column=1).value for r in range(1, len(texts) + 1)] == texts
    assert lists.cell(row=len(texts) + 1, column=1).value is None  # сплошной диапазон
    assert [lists.cell(row=r, column=2).value for r in range(1, 7)] == list(SOIL_TYPES)
    assert [lists.cell(row=r, column=3).value for r in range(1, 5)] == list(INSULATION_TYPES)
    objects = wb["Объекты"]
    col_a = [c.value for c in objects["A"]]
    assert "Объекты новой легенды" in col_a and "Объекты старого словаря" in col_a
    assert set(texts) <= set(col_a)
    assert len(objects._images) >= len(OBJECT_TYPES)


def test_data_validation(book):
    ws = load_workbook(book)["Result"]
    dvs = _validations(ws)
    comment_col = "H"  # «Комментарий» - 8-я колонка
    assert ws[f"{comment_col}1"].value == "Комментарий"
    cdv = dvs[f"{comment_col}2:{comment_col}10000"]
    assert cdv.type == "list" and not cdv.showErrorMessage
    assert cdv.showInputMessage and cdv.promptTitle and "через точку с запятой" in cdv.prompt
    assert cdv.formula1 == f"'Списки'!$A$1:$A${len(list_entries())}"
    soil = dvs["L2:L10000"]
    ins = dvs["R2:R10000"]
    assert ws["L1"].value == "Тип грунта" and ws["R1"].value == "Тип изоляции"
    assert soil.showErrorMessage and soil.errorStyle == "stop" and soil.error
    assert ins.showErrorMessage and ins.errorStyle == "stop" and ins.error
    assert soil.formula1 == f"'Списки'!$B$1:$B${len(SOIL_TYPES)}"
    assert ins.formula1 == f"'Списки'!$C$1:$C${len(INSULATION_TYPES)}"
    assert len(dvs) == 3


def test_filled_template_is_readable(tmp_path):
    path = tmp_path / "filled.xlsx"
    build_template(path, "Result")
    wb = load_workbook(path)
    ws = wb["Result"]
    notes = {2: "кип, КИП-14", 4: "река начало, р. Упа", 6: "река конец", 8: "вл; кран"}
    for i in range(2, 12):
        d = (i - 2) * 10
        ws.cell(row=i, column=1, value=d)
        ws.cell(row=i, column=2, value=-1.2)
        ws.cell(row=i, column=3, value=-1.0)
        if i in notes:
            ws.cell(row=i, column=8, value=notes[i])
    ws["J2"], ws["K2"], ws["P2"] = 0, 50, "торф"
    ws["J3"], ws["K3"], ws["P3"] = 50, 30, "Песчаный"
    ws["Q2"], ws["R2"] = 0, "битумное"
    ws["Q3"], ws["R3"] = 60, "пленочное полимерное"
    wb.save(path)
    survey, report = read_workbook(path)
    assert survey is not None and report.ok, report.format_text()
    notes_result = parse_comments(survey.comments.items)
    assert notes_result.unknown == 0 and notes_result.unpaired == 0
    assert [o.type_key for o in notes_result.objects] == ["kip", "water", "vl", "crane"]


def test_reopen_and_cli(tmp_path):
    out = tmp_path / "sub" / "cli.xlsx"
    assert main(["template", str(out), "--sheet-name", "Template"]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "Template"
    wb.save(tmp_path / "again.xlsx")
    assert load_workbook(tmp_path / "again.xlsx").sheetnames == wb.sheetnames
    with pytest.raises(SystemExit):
        main(["template", str(out), "--sheet-name", "Other"])
