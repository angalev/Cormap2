"""Модуль заказчика `cormap/refinery.py`: интерфейс из ТЗ 4.5 на настоящем модуле."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

# cormap/refinery.py — модуль заказчика, в репозиторий не входит (см. README,
# «Что не входит в репозиторий»). Без него этот файл целиком пропускается
# вместо ошибки сбора тестов.
refinery = pytest.importorskip("cormap.refinery")
from cormap.preprocess import RefineOptions, run_refinery_step

SOURCE = Path(__file__).resolve().parent.parent / "reference" / "test_data.xlsx"


@pytest.fixture
def book(tmp_path: Path) -> Path:
    path = tmp_path / "book.xlsx"
    shutil.copyfile(SOURCE, path)
    return path


def _result_rows(path: Path) -> list[tuple]:
    wb = load_workbook(path, data_only=True)
    rows = list(wb["Result"].iter_rows(values_only=True))
    wb.close()
    return rows


def _template_rows(path: Path) -> list[tuple]:
    wb = load_workbook(path, data_only=True)
    rows = list(wb["Template"].iter_rows(values_only=True))
    wb.close()
    return rows


def test_run_refinery_writes_result_and_keeps_source(book: Path, tmp_path: Path) -> None:
    before = book.read_bytes()
    out = tmp_path / "out.xlsx"
    settings = refinery.RefinerySettings(utz=True, utzdc=True, step=5, random_seed=1)
    result = refinery.run_refinery(book, settings, out)

    assert result.success, result.error_message
    assert result.output_file == out
    assert book.read_bytes() == before  # исходная книга нетронута
    source = _template_rows(book)
    rows = _result_rows(out)
    assert rows[0][: len(source[0])] == source[0]  # заголовки те же
    assert result.stats.input_rows == len(source) - 1
    assert result.stats.output_rows == len(rows) - 1
    assert result.stats.columns_generated == ["utz", "utzdc"]
    first, last = result.stats.distance_range_m
    assert (rows[1][0], rows[-1][0]) == (first, last)
    assert all(b[0] - a[0] <= 5 for a, b in zip(rows[1:], rows[2:], strict=False))
    assert all(r[1] is not None and r[2] is not None for r in rows[1:])  # U т-з вкл и выкл
    assert all(r[3] is None for r in rows[1:])  # флаг градиента не выставлен


def test_extra_columns_are_copied(book: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    result = refinery.run_refinery(book, refinery.RefinerySettings(utz=True, random_seed=1), out)
    assert result.success, result.error_message
    source = _template_rows(book)
    rows = _result_rows(out)
    width = len(source[0])
    for src, res in zip(source[1:], rows[1:], strict=False):
        assert res[9:width] == src[9:width]


def test_right_gradient_is_not_taken_from_left(book: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    settings = refinery.RefinerySettings(grru=True, grrd=True, random_seed=1)
    result = refinery.run_refinery(book, settings, out)  # левый градиент не выбран
    assert result.success, result.error_message
    source = _template_rows(book)
    rows = _result_rows(out)
    assert (rows[1][5], rows[1][6]) == (source[1][5], source[1][6])
    assert all(r[5] is not None and r[6] is not None for r in rows[1:])
    assert all(r[3] is None and r[4] is None for r in rows[1:])


def test_same_seed_gives_same_result(book: Path, tmp_path: Path) -> None:
    settings = refinery.RefinerySettings(utz=True, grluc=True, grldc=True, step=2, random_seed=42)
    one, two = tmp_path / "one.xlsx", tmp_path / "two.xlsx"
    assert refinery.run_refinery(book, settings, one).success
    assert refinery.run_refinery(book, settings, two).success
    assert _result_rows(one) == _result_rows(two)


@pytest.mark.parametrize(
    "flags",
    [
        {"grrd": True},  # было деление на ноль на коротких участках
        {"utz": True, "grluc": True, "grru": True, "grrd": True},  # был пустой диапазон
        {"utz": True, "utzd": True, "grlu": True, "grld": True, "grru": True, "grrd": True},
    ],
)
@pytest.mark.parametrize("step", [1, 5, 10])
def test_flag_sets_do_not_fail(book: Path, tmp_path: Path, flags: dict, step: int) -> None:
    for seed in range(2):
        settings = refinery.RefinerySettings(**flags, step=step, random_seed=seed)
        result = refinery.run_refinery(book, settings, tmp_path / "out.xlsx")
        assert result.success, f"{result.error_code}: {result.error_message}"


def test_errors_are_returned_not_raised(book: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    bad = refinery.run_refinery(book, refinery.RefinerySettings(step=5), out)
    assert (bad.success, bad.error_code, bad.stats) == (False, "bad_settings", None)

    wb = load_workbook(book)
    wb["Template"].title = "Данные"
    wb.save(book)
    result = refinery.run_refinery(book, refinery.RefinerySettings(utz=True), out)
    assert (result.success, result.error_code) == (False, "no_template")
    assert not out.exists()

    missing = refinery.run_refinery(tmp_path / "нет.xlsx", refinery.RefinerySettings(utz=True), out)
    assert (missing.success, missing.error_code) == (False, "read_error")


def test_service_step_calls_real_module(book: Path, tmp_path: Path, monkeypatch) -> None:
    # общая фикстура прячет настоящий модуль - здесь он возвращается на место
    monkeypatch.setitem(sys.modules, "cormap.refinery", refinery)
    out = tmp_path / "processed.xlsx"
    outcome = run_refinery_step(book, out, RefineOptions(utz=True, utzdc=True, random_seed=3))
    assert outcome.ok, outcome.error_message
    assert out.is_file()
    assert outcome.stats.output_rows > outcome.stats.input_rows


def test_service_step_survives_book_with_empty_tail(book: Path, tmp_path: Path, monkeypatch):
    """Случай с сервера: пустая оформленная ячейка в последней строке листа (1 048 576).

    Без обрезки модуль перебирает весь миллион строк и расходует больше гигабайта.
    """
    import tracemalloc

    from openpyxl.styles import Font

    monkeypatch.setitem(sys.modules, "cormap.refinery", refinery)
    wb = load_workbook(book)
    wb["Template"].cell(1_048_576, 13).font = Font(bold=True)
    wb.save(book)
    clean = run_refinery_step(book, tmp_path / "a.xlsx", RefineOptions(utz=True, random_seed=3))
    tracemalloc.start()
    try:
        outcome = run_refinery_step(
            book, tmp_path / "b.xlsx", RefineOptions(utz=True, random_seed=3)
        )
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert outcome.ok, outcome.error_message
    assert outcome.stats.input_rows == clean.stats.input_rows
    assert peak < 150 * 1024 * 1024, f"пик памяти {peak / 1e6:.0f} МБ"
