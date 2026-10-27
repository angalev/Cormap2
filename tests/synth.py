"""Генератор синтетической книги шаблона версии 2 (используется в тестах всех этапов)."""

from __future__ import annotations

import random
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook

HEADERS_V2 = [
    "Дистанция, м",
    "U т-з, В",
    "U т-з выкл, В",
    "Гр-нт лев. вкл, мВ",
    "Гр-нт лев. выкл, мВ",
    "Гр-нт прав. вкл, мВ",
    "Гр-нт прав. выкл, мВ",
    "Комментарий",
    "GPS",
    "Дистанция грун, м",
    "Уд. соп-е грунтов, Ом*м",
    "Тип грунта",
    "Дистанция глубины, м",
    "Глубина заложения трубы, м",
    "Дистанция повреждений, м",
    "Повреждения ИП",
    "Дистанция изоляции, м",
    "Тип изоляции",
]

SOILS = ["песчаный", "скальный", "суглинок", "супесь", "торф", "илистый"]
INSULATIONS = ["битумное", "заводское полиэтиленовое", "лакокрасочное", "пленочное полимерное"]
# Комментарии, каждый цикл содержит корректные пары «начало/конец».
NOTES = ["КИП, КИП-1", "УКЗ, УКЗ-3", "река начало, р. Упа", "река конец", "кран", "ВЛ-110 кВ"]


def make_workbook(
    path: str | Path,
    length_m: float = 50_000,
    step_m: float = 2.0,
    seed: int = 1,
    sheet: str = "Result",
    notes: Sequence[tuple[float, str]] | None = None,
) -> Path:
    """Создаёт книгу шаблона версии 2 с правдоподобными значениями во всех блоках.

    Блоки независимы: у каждого свой шаг дистанции. Результат детерминирован (seed).
    notes - список (дистанция, текст комментария): заменяет стандартный набор
    комментариев (пустой список - комментариев нет). Дистанция округляется до
    ближайшей точки сетки; несколько пометок в одной точке склеиваются через «; ».
    """
    custom: dict[int, str] | None = None
    if notes is not None:
        custom = {}
        for dist, text in notes:
            idx = round(dist / step_m)
            if not 0 <= idx <= int(length_m // step_m):
                raise ValueError(f"Дистанция комментария {dist} вне трассы 0..{length_m}")
            custom[idx] = f"{custom[idx]}; {text}" if idx in custom else text
    rnd = random.Random(seed)
    n = int(length_m // step_m) + 1

    def at(step: float, count: int) -> list[float]:
        return [round(i * step, 3) for i in range(count) if i * step <= length_m]

    soil_d = at(50.0, int(length_m // 50) + 1)
    depth_d = at(20.0, int(length_m // 20) + 1)
    ins_d = at(5000.0, int(length_m // 5000) + 1)
    soil_type_every = 40  # отметка типа грунта каждые 40 точек УЭС (2 км)
    dmg_d = at(max(length_m / 7, step_m), 7)
    note_every = max(int(500 // step_m), 1)  # объект каждые 500 м

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet)
    ws.append(HEADERS_V2)
    u = -1.5
    note_i = 0
    for i in range(n):
        u = max(-2.6, min(-0.7, u + rnd.uniform(-0.01, 0.01)))
        row: list[object] = [None] * len(HEADERS_V2)
        row[0] = round(i * step_m, 3)
        row[1] = round(u, 4)
        row[2] = round(u + 0.3 + rnd.uniform(-0.02, 0.02), 4)
        row[3] = rnd.randint(2, 15)
        row[4] = rnd.randint(1, 10)
        row[5] = rnd.randint(2, 15)
        row[6] = rnd.randint(1, 10)
        if custom is not None:
            if i in custom:
                row[7] = custom[i]
        elif i % note_every == 0:
            row[7] = NOTES[note_i % len(NOTES)]
            row[8] = f"{55 + rnd.random():.4f};{37 + rnd.random():.4f}"
            note_i += 1
        if i < len(soil_d):
            row[9] = soil_d[i]
            row[10] = round(rnd.uniform(10, 120), 1)
            if i % soil_type_every == 0:
                row[11] = SOILS[(i // soil_type_every) % len(SOILS)]
        if i < len(depth_d):
            row[12] = depth_d[i]
            row[13] = round(rnd.uniform(0.9, 2.0), 2)
        if i < len(dmg_d):
            row[14] = round(dmg_d[i], 3)
            row[15] = rnd.randint(1, 3)
        if i < len(ins_d):
            row[16] = ins_d[i]
            row[17] = INSULATIONS[i % len(INSULATIONS)]
        ws.append(row)
    wb.save(Path(path))
    return Path(path)


def make_sample_like(path: str | Path, seed: int = 7) -> Path:
    """Книга шаблона версии 2, имитирующая образец приложения № 11 (reference/obrazec_list*.png).

    5,5 км, шаг 2 м; U т-з около -1,5...-1,3 В с плавными ступеньками, U т-з выкл ближе к
    нулю на ~0,2 В; градиенты шумят (левые 10-35, правые 5-28 мВ); глубина - раз в ~400 м
    (0,8-1,4 м); УЭС - раз в 100 м (100-450 Ом*м); 3 типа грунта, 3 типа изоляции;
    комментарии: КИП с подписями, краны, дороги, пересечения, ВЛ, река, УКЗ, ПИП;
    блок «Повреждения» - три точки с классами 1, 2, 3. Результат детерминирован.
    """
    length, step = 5500.0, 2.0
    rnd = random.Random(seed)
    n = int(length // step) + 1

    # плавный профиль U т-з: (дистанция, значение), между точками - линейно
    profile = [
        (0, -1.50), (1000, -1.48), (1030, -1.62), (1500, -1.50), (2000, -1.33), (3000, -1.31),
        (3150, -1.38), (3400, -1.35), (4000, -1.38), (4200, -1.49), (4480, -1.47), (4500, -1.35),
        (5000, -1.30), (5400, -1.20), (5500, -1.10),
    ]  # fmt: skip

    def u_at(d: float) -> float:
        for (d0, v0), (d1, v1) in zip(profile, profile[1:], strict=False):
            if d0 <= d <= d1:
                return v0 + (v1 - v0) * (d - d0) / (d1 - d0)
        return profile[-1][1]

    notes: dict[int, str] = {}

    def note(dist: float, text: str) -> None:
        idx = round(dist / step)
        notes[idx] = f"{notes[idx]}; {text}" if idx in notes else text

    kips = [0, 150, 500, 1000, 1050, 1900, 2900, 3700, 4250, 4450, 4500, 4750, 5300, 5450]
    for i, d in enumerate(kips, 1):
        note(d, f"КИП, КИП-{i}")
    note(0, "кран")
    note(4460, "кран")
    note(1010, "асфальт а/д начало")
    note(1030, "асфальт а/д конец")
    note(5200, "грунт а/д начало")
    note(5215, "грунт а/д конец")
    note(3700, "пересечение с газопроводом")
    note(4700, "пересечение с кабелем")
    note(2400, "ВЛ-110 кВ")
    note(3000, "река начало, р. Упа")
    note(3150, "река конец")
    note(4260, "УКЗ, УКЗ-1")
    note(1800, "ПИП")
    note(3500, "ПИП")

    soil_marks = {0: "суглинок", 18: "песчаный", 39: "супесь"}  # индексы точек УЭС (шаг 100 м)
    ins_marks = {0.0: "битумное", 2000.0: "заводское полиэтиленовое", 4200.0: "лакокрасочное"}
    depth_d = [i * 400.0 for i in range(int(length // 400) + 1)] + [length]
    depth_v = [round(rnd.uniform(0.9, 1.4), 2) for _ in depth_d]
    depth_v[7] = 0.80  # одна точка ровно на нормативе
    depth_v[8] = 0.81
    soil_d = [i * 100.0 for i in range(int(length // 100) + 1)]
    dmg = [(800.0, 1), (2600.0, 2), (4000.0, 3)]

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Result")
    ws.append(HEADERS_V2)
    for i in range(n):
        d = round(i * step, 3)
        row: list[object] = [None] * len(HEADERS_V2)
        u = u_at(d) + rnd.uniform(-0.012, 0.012)
        row[0] = d
        row[1] = round(u, 4)
        row[2] = round(u + 0.2 + rnd.uniform(-0.03, 0.03), 4)
        row[3] = rnd.randint(10, 35)
        row[4] = rnd.randint(8, 32)
        row[5] = rnd.randint(5, 28)
        row[6] = rnd.randint(4, 26)
        if i in notes:
            row[7] = notes[i]
        if i < len(soil_d):
            row[9] = soil_d[i]
            row[10] = rnd.randint(100, 450)
            row[11] = soil_marks.get(i)
        if i < len(depth_d):
            row[12] = depth_d[i]
            row[13] = depth_v[i]
        if i < len(dmg):
            row[14], row[15] = dmg[i]
        if i < len(ins_marks):
            dist_i = sorted(ins_marks)[i]
            row[16], row[17] = dist_i, ins_marks[dist_i]
        ws.append(row)
    wb.save(Path(path))
    return Path(path)
