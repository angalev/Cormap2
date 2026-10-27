"""Предварительная обработка: настройки, правила ТЗ 4.5, вызов модуля (на заглушке), CLI."""

from __future__ import annotations

import json
import sys

import pytest
from openpyxl import load_workbook

from cormap.cli import main
from cormap.preprocess import (
    FLAG_NAMES,
    RefineOptions,
    refinery_available,
    run_refinery_step,
    validate,
)
from tests.synth import make_workbook


def fields_of(options: RefineOptions) -> list[str]:
    return [f for f, _ in validate(options)]


# --- настройки ---------------------------------------------------------------------------


def test_options_dict_roundtrip():
    options = RefineOptions(utz=True, grlu=True, step=7, random_seed=123)
    data = options.to_dict()
    assert json.loads(json.dumps(data)) == data
    assert RefineOptions.from_dict(data) == options
    assert set(FLAG_NAMES) <= set(data) and len(FLAG_NAMES) == 11
    assert RefineOptions.from_dict({"utz": True, "лишнее": 1}).utz is True
    assert RefineOptions.from_dict(None) == RefineOptions()


# --- 7 правил: нарушение и допустимый случай ---------------------------------------------------


def test_no_flags_is_allowed_and_means_skip():
    assert validate(RefineOptions()) == []
    assert RefineOptions().any_flag is False
    assert validate(RefineOptions(utz=True)) == []
    assert RefineOptions(utz=True).any_flag is True


def test_rule2_utzdc_only_with_utz_without_utzd():
    assert fields_of(RefineOptions(utzdc=True)) == ["utzdc"]
    assert fields_of(RefineOptions(utz=True, utzd=True, utzdc=True)) == ["utzdc"]
    assert validate(RefineOptions(utz=True, utzdc=True)) == []


def test_rule3_left_pairs():
    assert fields_of(RefineOptions(grlu=True, grluc=True)) == ["grluc"]
    assert "grldc" in fields_of(RefineOptions(grlu=True, grld=True, grldc=True))
    assert validate(RefineOptions(grluc=True, grld=True)) == []


def test_rule4_grldc_needs_left_on():
    assert fields_of(RefineOptions(grldc=True)) == ["grldc"]
    assert validate(RefineOptions(grlu=True, grldc=True)) == []
    assert validate(RefineOptions(grluc=True, grldc=True)) == []


def test_rule5_right_pairs():
    assert fields_of(RefineOptions(grru=True, grruc=True)) == ["grruc"]
    assert "grrdc" in fields_of(RefineOptions(grru=True, grrd=True, grrdc=True))
    assert validate(RefineOptions(grruc=True, grrd=True)) == []


def test_rule6_grrdc_needs_right_on():
    assert fields_of(RefineOptions(grrdc=True)) == ["grrdc"]
    assert validate(RefineOptions(grru=True, grrdc=True)) == []
    assert validate(RefineOptions(grruc=True, grrdc=True)) == []


@pytest.mark.parametrize("step", [0, 11, -1, 2.5, "5", True])
def test_rule7_bad_step(step):
    assert fields_of(RefineOptions(utz=True, step=step)) == ["step"]


@pytest.mark.parametrize("step", [1, 5, 10])
def test_rule7_good_step(step):
    assert validate(RefineOptions(utz=True, step=step)) == []


# --- модуль не установлен ------------------------------------------------------------------------


def test_not_installed(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "refinery", None)  # импорт выбросит ImportError
    assert refinery_available() is False
    out = run_refinery_step(tmp_path / "a.xlsx", tmp_path / "b.xlsx", RefineOptions(utz=True))
    assert not out.ok and out.error_code == "not_installed" and out.stats is None


def test_available(fake_refinery):
    assert refinery_available() is True


# --- вызов модуля ------------------------------------------------------------------------------


@pytest.fixture
def src(tmp_path):
    return make_workbook(tmp_path / "in.xlsx", length_m=200, step_m=10, sheet="Template")


def test_all_settings_passed(fake_refinery, src, tmp_path):
    options = RefineOptions(
        utz=True, utzdc=True, grluc=True, grldc=True, grru=True, grrdc=True, step=3, random_seed=42
    )
    dst = tmp_path / "sub" / "out.xlsx"
    out = run_refinery_step(src, dst, options)
    assert out.ok and dst.is_file()
    (call,) = fake_refinery.calls
    assert call["input_file"] == src and call["output_file"] == dst
    got = call["settings"]
    for name in FLAG_NAMES:
        assert getattr(got, name) is getattr(options, name), name
    assert got.step == 3 and got.random_seed == 42
    assert out.stats.input_rows == 21 and out.stats.output_rows == 21
    assert out.stats.distance_range_m == (0.0, 200.0)
    data = out.stats.to_dict()
    assert json.loads(json.dumps(data)) == data  # сериализуется в JSON (кортеж -> список)
    ws = load_workbook(dst)["Result"]
    assert ws.max_row == 22


def test_size_limits_checked_before_module(fake_refinery, src, tmp_path, monkeypatch):
    """Модуль грузит книгу в память целиком: слишком большая до него не доходит."""
    monkeypatch.setattr("cormap.reader.MAX_UNPACKED_BYTES", 1000)  # книга распакована больше
    dst = tmp_path / "out.xlsx"
    out = run_refinery_step(src, dst, RefineOptions(utz=True))
    assert not out.ok and out.error_code == "file_too_big"
    assert fake_refinery.calls == [] and not dst.exists()


def test_not_a_book_does_not_reach_module(fake_refinery, tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    out = run_refinery_step(bad, tmp_path / "out.xlsx", RefineOptions(utz=True))
    assert not out.ok and out.error_code == "bad_file"
    assert fake_refinery.calls == []


def _add_stray_cell(path, row=1_048_576):
    """Пустая оформленная ячейка далеко внизу листа: так книгу иногда сохраняет Excel."""
    from openpyxl.styles import Font

    wb = load_workbook(path)
    wb["Template"].cell(row, 13).font = Font(bold=True)
    wb.save(path)
    wb.close()


def test_empty_tail_is_cut_before_module(fake_refinery, src, tmp_path, monkeypatch):
    """Лист объявлен на миллион строк, данных 21: модуль получает книгу без хвоста."""
    _add_stray_cell(src)
    seen = {}
    original = sys.modules["refinery"].run_refinery

    def spy(input_file, settings, output_file=None):
        wb = load_workbook(input_file, read_only=True)
        seen["rows"], seen["name"] = wb["Template"].max_row, input_file.name
        wb.close()
        return original(input_file, settings, output_file)

    monkeypatch.setattr(sys.modules["refinery"], "run_refinery", spy)
    dst = tmp_path / "out.xlsx"
    out = run_refinery_step(src, dst, RefineOptions(utz=True))
    assert out.ok, out.error_message
    assert seen == {"rows": 22, "name": "out_trimmed.xlsx"}
    assert out.stats.input_rows == 21
    assert sorted(p.name for p in tmp_path.iterdir()) == ["in.xlsx", "out.xlsx"]  # копия убрана
    source = load_workbook(src, read_only=True)
    assert source["Template"].max_row == 1_048_576  # исходная книга не тронута
    source.close()


def test_empty_tail_copy_removed_on_module_error(fake_refinery, src, tmp_path):
    _add_stray_cell(src)
    fake_refinery.mode = "raise"
    out = run_refinery_step(src, tmp_path / "out.xlsx", RefineOptions(utz=True))
    assert not out.ok and out.error_code == "calculation_error"
    assert [p.name for p in tmp_path.iterdir()] == ["in.xlsx"]


def test_too_many_real_rows_do_not_reach_module(fake_refinery, src, tmp_path, monkeypatch):
    monkeypatch.setattr("cormap.preprocess.MAX_DATA_ROWS", 10)  # в книге 21 строка данных
    out = run_refinery_step(src, tmp_path / "out.xlsx", RefineOptions(utz=True))
    assert not out.ok and out.error_code == "too_many_rows"
    assert fake_refinery.calls == [] and [p.name for p in tmp_path.iterdir()] == ["in.xlsx"]


def test_module_failure(fake_refinery, src, tmp_path):
    fake_refinery.mode = "fail"
    dst = tmp_path / "out.xlsx"
    out = run_refinery_step(src, dst, RefineOptions(utz=True))
    assert not out.ok and out.error_code == "bad_input"
    assert "Нет данных" in out.error_message and out.stats is None
    assert not dst.exists()


def test_module_exception(fake_refinery, src, tmp_path, caplog):
    fake_refinery.mode = "raise"
    out = run_refinery_step(src, tmp_path / "out.xlsx", RefineOptions(utz=True))
    assert not out.ok and out.error_code == "calculation_error"
    assert "внезапный сбой" in out.error_message and "Traceback" not in out.error_message
    assert any(r.exc_info for r in caplog.records)  # трассировка ушла в журнал


def test_success_without_file(fake_refinery, src, tmp_path):
    fake_refinery.mode = "no_file"
    out = run_refinery_step(src, tmp_path / "out.xlsx", RefineOptions(utz=True))
    assert not out.ok and out.error_code == "write_error"


def test_success_without_stats(fake_refinery, src, tmp_path, monkeypatch):
    def run(input_file, settings, output_file=None):
        output_file.write_bytes(b"x")
        return type("R", (), {"success": True, "stats": None})()

    monkeypatch.setattr(sys.modules["refinery"], "run_refinery", run)
    out = run_refinery_step(src, tmp_path / "out.xlsx", RefineOptions(utz=True))
    assert out.ok and out.stats is None


# --- командная строка --------------------------------------------------------------------------


def test_cli_refine_ok(fake_refinery, src, tmp_path, capsys):
    dst = tmp_path / "cli.xlsx"
    code = main(["refine", str(src), str(dst), "--utz", "--grlu", "--step", "4", "--seed", "9"])
    assert code == 0 and dst.is_file()
    settings = fake_refinery.calls[0]["settings"]
    assert settings.utz and settings.grlu and not settings.utzd
    assert settings.step == 4 and settings.random_seed == 9
    out = capsys.readouterr().out
    assert "Строк на входе: 21" in out


def test_cli_refine_rules_before_call(fake_refinery, src, tmp_path, capsys):
    code = main(["refine", str(src), str(tmp_path / "x.xlsx"), "--utzdc"])
    assert code == 2 and fake_refinery.calls == []
    assert "utzdc" in capsys.readouterr().err
    assert main(["refine", str(src), str(tmp_path / "x.xlsx")]) == 2  # ни одного флажка


def test_cli_refine_module_error(fake_refinery, src, tmp_path, capsys):
    fake_refinery.mode = "fail"
    code = main(["refine", str(src), str(tmp_path / "x.xlsx"), "--utz"])
    assert code == 1
    assert "[bad_input]" in capsys.readouterr().err


def test_cli_refine_not_installed(src, tmp_path, capsys, monkeypatch):
    monkeypatch.setitem(sys.modules, "refinery", None)
    assert main(["refine", str(src), str(tmp_path / "x.xlsx"), "--utz"]) == 1
    assert "[not_installed]" in capsys.readouterr().err
